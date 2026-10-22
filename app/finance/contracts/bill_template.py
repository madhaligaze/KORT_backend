"""Шаблон счетов Excel: скачать в разделе «Счета», заполнить, загрузить обратно.

Финансисты BBC выставляют счета по своему шаблону Excel («Шаблон Ексель.xlsx»,
22.10.2026): одна строка - один счёт. Колонки и их порядок - как у них: «Наша
Фирма», «Наш Сотрудник», «Заказчик», «По 1C», «№ Договора», «Дата заключения
Договора», «Период с/по», «Сумма Договора» - и жёлтые «№ Счета» и «Дата
выстав. Счета», которые заполняет человек.

В файл идут те же счета, что бот прислал в канал, в том же порядке
(`bills_bot.listed`). Листы (порядок задал владелец 23.10.2026):
* «Общий реестр» - все счета вперемешку;
* «Аренда», «Абонентские услуги» (абонентская плата) - те же счета по виду;
* по листу на каждое юрлицо («Наша Фирма») - здесь счета и выставляют:
  жёлтые клетки только тут.
«№ Счета» и дату общие листы берут с листа юрлица формулой - по скрытому коду
счёта, а не по адресу клетки: отсортированный лист юрлица не подсунет общему
реестру чужой номер. Одно и то же дважды заполнять не нужно.

Загруженный файл в книги сам не пишет: KORT разбирает его в строки, человек
видит, что из него вышло, и каждую строку в книгу её юрлица записывает та же
запись, что у карточки счёта (`POST /bills/{id}` → `bill_book.save`), со всеми
её проверками.

Разбор - по названиям колонок, а не по буквам: шаблон правят и вставляют в
него колонки. Строка счёта узнаётся по скрытому «Код KORT», а в шаблоне
финансистов без кода - по юрлицу, клиенту, № договора и «Периоду с». Счёт на
нескольких листах - одна строка, если везде вписано одно и то же; разное - отказ
этому счёту, а не выбор наугад. Не нашлась, нашлось две, дважды на одном листе -
тоже отказ строке словами; остальные строки файла от этого не страдают.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Callable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.finance.contracts.bill_book import BookError, wanted_values
from app.finance.contracts.bills import DONE, Bill, Board, read_date
from app.finance.layout import Column, LayoutError, norm, resolve_layout

#: Сколько весит загружаемый шаблон: сотня строк - это десятки килобайт.
MAX_BYTES = 2 * 1024 * 1024
#: Сколько он весит распакованным: .xlsx - zip, и маленький файл может раскрыться в гигабайты.
MAX_UNPACKED = 50 * 1024 * 1024
#: Сколько строк просматривать под шапкой и где искать саму шапку.
MAX_ROWS = 1000
HEADER_SEARCH = 15
EPOCH = date(1899, 12, 30)

#: Жёлтые колонки - их заполняет человек (как в шаблоне финансистов).
#: Цвет - с непрозрачной альфой (FF…), как в файле финансистов: «FFF2CC»
#: openpyxl пишет как 00FFF2CC, и LibreOffice показал бы клетку без заливки.
FILL = PatternFill(fill_type="solid", fgColor="FFFFF2CC")
#: Только строчными. «DD.MM.YYYY» Excel читает как дату, а просмотр Apple
#: (айфон, 23.10.2026) - по правилам Unicode, где «D» - номер дня в году:
#: 23.10.2026 показывался «296.10.2026», 23.02 - «54.02».
DATE_FORMAT = "dd.mm.yyyy"
#: Имена листов выбраны так, чтобы Excel рисовал ярлыки на одной линии: русское
#: имя без букв с хвостом вниз (р, у, д) он ставит на пару пикселей ниже
#: соседей - «Счета» и «Абонентская плата» выглядели «на разных уровнях»
#: (владелец, 23.10.2026; так же и в книге, созданной самим Excel). Имена -
#: выбор владельца.
SHEET_TITLE = "Общий реестр"
RENT_SHEET = "Аренда"
ABON_SHEET = "Абонентские услуги"
#: Общие листы: «№ Счета» и дату на них тянет формула с листа юрлица. Прежние
#: имена - для шаблонов, скачанных до 23.10.2026.
SUMMARY_SHEETS = frozenset({SHEET_TITLE, RENT_SHEET, ABON_SHEET, "Счета", "Абонентская плата"})
#: Ячейки, по которым формула общего листа ищет счёт на листе юрлица.
LOOKUP_ROWS = 2000

#: Колонки шаблона: буква, заголовок, ширина - как у финансистов; `field` -
#: что из счёта туда кладётся.
LAYOUT: tuple[tuple[str, str, float, str], ...] = (
    ("B", "Наша\nФирма", 9.3, "own"),
    ("C", "Наш\nСотрудник", 13.0, "staff"),
    ("D", "Заказчик\n(Название Фирмы)", 24.4, "customer"),
    ("E", "По 1C", 22.0, "onec_name"),
    ("F", "№\nДоговора", 10.6, "number"),
    ("G", "Дата заключения \nДоговора", 17.9, "signed_at"),
    ("I", "Период\nс (нач.)", 10.6, "period_from"),
    ("J", "Период\nпо (зав.)", 10.6, "period_to"),
    ("L", "Сумма\nДоговора", 14.6, "amount"),
    ("N", "№ \nСчета", 12.0, "bill_no"),
    ("O", "Дата выстав.\nСчета", 16.3, "bill_date"),
)
#: Узкие колонки-просветы между блоками - как в шаблоне финансистов.
GAPS = {"A": 2.7, "H": 4.0, "K": 3.0, "M": 2.9}
#: Скрытая колонка - по ней загруженный шаблон узнаёт свои счета.
KORT_COLUMN = "Q"
KORT_TITLE = "Код KORT"

#: Колонки загружаемого шаблона - только точные названия: из него пишут в книги.
COLUMNS: tuple[Column, ...] = (
    Column("own", "Наша Фирма", ("наша фирма",), strict=True),
    Column("customer", "Заказчик (Название Фирмы)", ("заказчик (название фирмы)", "заказчик"), strict=True),
    Column("number", "№ Договора", ("№ договора",), strict=True),
    Column("period_from", "Период с (нач.)", ("период с (нач.)", "период с"), required=False, strict=True),
    Column("bill_no", "№ Счета", ("№ счета",), strict=True),
    Column("bill_date", "Дата выстав. Счета", ("дата выстав. счета", "дата выставления счета"), strict=True),
    Column("bill_period", "Счет за Период (Коммент)", ("счет за период (коммент)",), required=False, strict=True),
    Column("note", "Доп. Коммент.", ("доп. коммент.",), required=False, strict=True),
    Column("kort", KORT_TITLE, ("код kort",), required=False, strict=True),
)

#: Что со строкой файла: записать, уже в книге, не заполнена, отказ.
READY = "ready"
SAME = "same"
EMPTY = "empty"
ERROR = "error"
STATUSES = (READY, SAME, EMPTY, ERROR)


class TemplateError(Exception):
    """Шаблон не разобрать целиком - с текстом для человека."""


# ── Скачать ──────────────────────────────────────────────────────────────────


def _cell_value(bill: Bill, name: str) -> Any:
    if name in ("period_from", "period_to"):
        return getattr(bill, name)
    if name == "signed_at":
        return read_date(bill.signed_at) or (bill.signed_at or None)
    if name == "amount":
        if bill.amount is None:
            return None
        return int(bill.amount) if bill.amount == bill.amount.to_integral() else float(bill.amount)
    if name in ("bill_no", "bill_date"):
        return None  # заполняет человек: в шаблон идут невыставленные
    return getattr(bill, name) or None


def _quoted(title: str) -> str:
    """Имя листа в формуле: «'BBC Omar'»."""
    return "'" + title.replace("'", "''") + "'"


def _pull(source: str, column: str, row: int) -> str:
    """Формула общего листа: значение колонки с листа юрлица - по коду счёта, а не по адресу.

    Пусто на листе юрлица - пусто и здесь (`INDEX` пустой клетки дал бы 0, а
    в колонке даты - 00.01.1900); счёта там нет - тоже пусто, не #N/A.
    """
    sheet = _quoted(source)
    last = LOOKUP_ROWS + 1
    look = f"INDEX({sheet}!${column}$2:${column}${last},MATCH(${KORT_COLUMN}{row},{sheet}!${KORT_COLUMN}$2:${KORT_COLUMN}${last},0))"
    return f'=IFERROR(IF({look}="","",{look}),"")'


def _firm_titles(items: Sequence[Bill]) -> dict[str, str]:
    """Юрлицо → имя его листа: по алфавиту, без знаков, которых не терпит Excel, не длиннее 31 и без повторов."""
    used = {title.casefold() for title in SUMMARY_SHEETS}
    out: dict[str, str] = {}
    for own in sorted({bill.own for bill in items}, key=lambda name: (name.casefold(), name)):
        base = re.sub(r"[\[\]:*?/\\]", "-", own).strip().strip("'")[:31].strip() or "Без юрлица"
        title, number = base, 2
        while title.casefold() in used:
            suffix = f" ({number})"
            title, number = base[: 31 - len(suffix)] + suffix, number + 1
        used.add(title.casefold())
        out[own] = title
    return out


def _fill(sheet: Any, items: Sequence[Bill], sources: dict[str, str] | None) -> None:
    """Лист шаблона. `sources` - юрлицо → лист, откуда тянуть № и дату; `None` - их вписывают здесь."""
    bold = Font(bold=True)
    head = Alignment(wrap_text=True, vertical="center")
    for letter, title, width, _ in LAYOUT:
        top = sheet[f"{letter}1"]
        top.value, top.font, top.alignment = title, bold, head
        sheet.column_dimensions[letter].width = width
    for letter, width in GAPS.items():
        sheet.column_dimensions[letter].width = width
    sheet.row_dimensions[1].height = 32
    if sources is None:
        sheet["N1"].fill = FILL
        sheet["O1"].fill = FILL
    else:
        for letter in ("N", "O"):
            sheet[f"{letter}1"].comment = Comment("Заполняется само с листа юрлица - вписывайте там", "KORT")
    # Код счёта - в скрытой колонке: по нему KORT узнаёт свою строку, даже если
    # название клиента в файле поправили, а общий лист - строку на листе юрлица.
    sheet[f"{KORT_COLUMN}1"].value = KORT_TITLE
    sheet.column_dimensions[KORT_COLUMN].hidden = True
    last = 1
    for last, bill in enumerate(items, start=2):
        sheet[f"{KORT_COLUMN}{last}"].value = bill.id
        for letter, _, _, name in LAYOUT:
            cell = sheet[f"{letter}{last}"]
            if name in ("bill_no", "bill_date"):
                if sources is None:
                    cell.fill = FILL
                else:
                    cell.value = _pull(sources[bill.own], letter, last)
            else:
                cell.value = _cell_value(bill, name)
            if name in ("signed_at", "period_from", "period_to", "bill_date"):
                cell.number_format = DATE_FORMAT
            if name == "amount":
                cell.number_format = "#,##0"
    sheet.freeze_panes = "A2"
    if last > 1:
        rule = DataValidation(
            type="date",
            operator="greaterThan",
            formula1="DATE(2020,1,1)",
            showErrorMessage=True,
            errorTitle="Дата выстав. Счета",
            error="Впишите дату счёта: 22.10.2026",
        )
        rule.add(f"O2:O{last}")
        sheet.add_data_validation(rule)


def build(items: Sequence[Bill], kind_of: Callable[[Bill], str | None] = lambda bill: None) -> bytes:
    """Шаблон: «Общий реестр», «Аренда», «Абонентские услуги» и лист на каждое юрлицо.

    `kind_of` - лист вида счёта (`RENT_SHEET`, `ABON_SHEET`) или `None`: «без
    вида услуги» есть только в общем реестре и на листе своего юрлица.
    """
    book = Workbook()
    # Формулы общих листов посчитать при открытии: своих значений у них в файле нет.
    book.calculation.fullCalcOnLoad = True
    firms = _firm_titles(items)
    summary = book.active
    summary.title = SHEET_TITLE
    _fill(summary, items, firms)
    # Порядок листов один и тот же каждый день: аренды нет - лист с одной шапкой.
    for title in (RENT_SHEET, ABON_SHEET):
        _fill(book.create_sheet(title), [bill for bill in items if kind_of(bill) == title], firms)
    for own, title in firms.items():
        _fill(book.create_sheet(title), [bill for bill in items if bill.own == own], None)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def filename(today: date) -> str:
    """«Счета к выставлению 22.10.2026.xlsx»."""
    return f"Счета к выставлению {today:%d.%m.%Y}.xlsx"


def period_text(bill: Bill) -> str:
    """«с 22.10.2026 года по 21.11.2026 года» - так «Счет за Период» пишут в книгах (как карточка)."""
    if bill.period_from is None:
        return ""
    text = f"с {bill.period_from:%d.%m.%Y} года"
    return f"{text} по {bill.period_to:%d.%m.%Y} года" if bill.period_to else text


# ── Загрузить ────────────────────────────────────────────────────────────────


def _open(data: bytes) -> Any:
    if not data:
        raise TemplateError("Файл пустой")
    if len(data) > MAX_BYTES:
        raise TemplateError("Файл больше 2 МБ - это не шаблон счетов")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED:
                raise TemplateError("Файл распаковывается в слишком большой объём - это не шаблон счетов")
    except zipfile.BadZipFile as exc:
        raise TemplateError("Это не файл .xlsx - сохраните шаблон в формате Excel (.xlsx)") from exc
    try:
        return load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl бросает много типов
        raise TemplateError("Файл не читается как .xlsx - сохраните шаблон в формате Excel (.xlsx)") from exc


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    return " ".join(str(value).split())


def _number(value: Any) -> str:
    """«№ Счета» для сравнения: без пробелов и регистра.

    В книгах у колонки формат «#,##0», и номер от тысячи книга показывает с
    неразрывным пробелом: 990101 - «990 101». В файле тот же номер - «990101».
    Сравнение по показанному тексту выдавало записанный счёт за «в книге уже
    другой номер» (сквозная проверка на копиях книг, 23.10.2026).
    """
    return "".join(_text(value).split()).casefold()


def _date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 36000 <= value <= 73000:
        return EPOCH + timedelta(days=int(value))  # дата числом в ячейке без формата даты
    text = str(value).strip()
    parsed = read_date(text)
    if parsed is not None:
        return parsed
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _find(sheet: Any) -> tuple[int, dict[str, int], list[tuple[Any, ...]]]:
    """Строка шапки, колонки по названиям и строки под шапкой."""
    rows = []
    for row in sheet.iter_rows(values_only=True, max_row=HEADER_SEARCH + MAX_ROWS):
        rows.append(tuple(row))
    for number, row in enumerate(rows[:HEADER_SEARCH]):
        header = [_text(cell) for cell in row]
        if not any(norm(cell) in ("№ счета",) for cell in header):
            continue
        try:
            layout = resolve_layout(sheet.title, COLUMNS, header)
        except LayoutError as exc:
            raise TemplateError(f"Шапку шаблона не разобрать: {exc}") from exc
        columns = {key: placement.index for key, placement in layout.placements.items() if placement.index is not None}
        return number, columns, rows[number + 1 :]
    raise TemplateError("В файле нет шапки шаблона - колонки «№ Счета» не нашлось")


@dataclass
class Row:
    """Строка файла: какой это счёт и что с ней делать."""

    sheet: str
    row: int
    own: str
    customer: str
    number: str
    bill: Bill | None = None
    status: str = READY
    #: Почему строку не записать (`ERROR`).
    problem: str = ""
    #: Что уже стоит в книге (`SAME`) - не отказ, а сведение.
    note: str = ""
    values: dict[str, str] = field(default_factory=dict)
    defaults: dict[str, str] = field(default_factory=dict)
    #: Что вписано в строку как есть: так сверяются листы одного счёта.
    entered: tuple[str, ...] = ()

    @property
    def filled(self) -> bool:
        return self.status != EMPTY

    def fail(self, problem: str) -> None:
        self.status, self.problem, self.note = ERROR, problem, ""
        self.values, self.defaults = {}, {}

    def out(self) -> dict[str, Any]:
        bill = self.bill
        return {
            "sheet": self.sheet,
            "row": self.row,
            "id": bill.id if bill else None,
            "own": bill.own if bill else self.own,
            "customer": bill.customer if bill else self.customer,
            "number": bill.number if bill else self.number,
            "month": bill.month if bill else "",
            "period_from": bill.period_from.isoformat() if bill and bill.period_from else None,
            "period_to": bill.period_to.isoformat() if bill and bill.period_to else None,
            "amount": str(bill.amount) if bill and bill.amount is not None else None,
            "status": self.status,
            "problem": self.problem,
            "note": self.note,
            "values": self.values,
            "defaults": self.defaults,
        }


def _match(
    own: str, customer: str, number: str, period: date | None, code: str, everything: list[Bill], by_id: dict[str, Bill]
) -> tuple[Bill | None, str]:
    """Счёт строки - по коду KORT или по договору и «Периоду с». Не один - отказ словами."""
    bill = by_id.get(code) if code else None
    if bill is not None:
        given = [(own, bill.own), (customer, bill.customer), (number, bill.number)]
        if any(text and norm(text) != norm(real) for text, real in given):
            return None, f"Строку правили: в ней другой договор, чем у её кода KORT ({bill.customer} {bill.number}). Скачайте шаблон заново"
        return bill, ""
    if not (customer and number):
        return None, "В строке не видно договора: нет «Заказчика» или «№ Договора»"
    candidates = [
        item
        for item in everything
        if norm(item.customer) == norm(customer)
        and norm(item.number) == norm(number)
        and (not own or norm(item.own) == norm(own))
    ]
    if period is not None:
        candidates = [item for item in candidates if item.period_from in (None, period)]
    where = ", ".join(part for part in (own, customer, number) if part)
    if not candidates:
        when = f" с периодом от {period:%d.%m.%Y}" if period else ""
        return None, f"В книге-сводке нет счёта {where}{when}"
    if len(candidates) > 1:
        if period is None:
            return None, f"У договора {where} в книге-сводке несколько месяцев - впишите «Период с», KORT не угадывает"
        return None, f"В книге-сводке несколько строк {where} за этот период - какую брать, KORT не угадывает"
    return candidates[0], ""


def _read_row(
    sheet: str, line: int, columns: dict[str, int], cells: tuple[Any, ...], everything: list[Bill], by_id: dict[str, Bill]
) -> Row | None:
    def value(key: str) -> Any:
        index = columns.get(key)
        return cells[index] if index is not None and index < len(cells) else None

    own, customer, number = _text(value("own")), _text(value("customer")), _text(value("number"))
    code = _text(value("kort"))
    bill_no, raw_date = _text(value("bill_no")), value("bill_date")
    if not any((own, customer, number, code, bill_no, _text(raw_date))):
        return None  # пустая строка листа
    row = Row(sheet=sheet, row=line, own=own, customer=customer, number=number)
    issued_on = _date(raw_date)
    row.entered = (
        _number(bill_no),
        issued_on.isoformat() if issued_on else _text(raw_date),
        _text(value("bill_period")),
        _text(value("note")),
    )
    bill, problem = _match(own, customer, number, _date(value("period_from")), code, everything, by_id)
    row.bill = bill
    if not bill_no and not _text(raw_date):
        row.status = EMPTY  # не заполнена - не ошибка: финансист выставил не всё
        return row
    if bill is None:
        row.fail(problem)
        return row
    issued = _date(raw_date)
    if not bill_no:
        row.fail("Не вписан «№ Счета»")
        return row
    if issued is None:
        raw = _text(raw_date)
        row.fail("«Дата выстав. Счета» - " + (f"не дата: «{raw}»" if raw else "пусто"))
        return row
    values = {"bill_no": bill_no, "bill_date": issued.isoformat()}
    for key in ("bill_period", "note"):
        text = _text(value(key))
        if text:
            values[key] = text  # пустая клетка файла не стирает то, что уже стоит в книге
    try:
        wanted_values(values)  # те же правила, что у записи: длина, «=», «+», «@», год
    except BookError as exc:
        row.fail(str(exc))
        return row
    if bill.group == DONE:
        if _number(bill.bill_no) == _number(bill_no):
            row.status = SAME
            row.note = f"В книге уже № {bill.bill_no}" + (f" от {bill.bill_date}" if bill.bill_date else "")
            return row
        row.fail(
            f"В книге уже № {bill.bill_no}" + (f" от {bill.bill_date}" if bill.bill_date else "")
            + " - KORT его не перезаписывает. Исправить - в карточке счёта"
        )
        return row
    row.values = values
    if "bill_period" not in values and (text := period_text(bill)):
        # Как в карточке: при выставлении «Счет за Период» - из «Периода с/по», если в книге он пуст.
        row.defaults = {"bill_period": text}
    return row


def _merge(rows: list[Row]) -> list[Row]:
    """Один счёт на нескольких листах - одна строка; разное на листах - отказ.

    Общие листы тянут № и дату с листа юрлица формулой, поэтому счёт стоит в
    файле два-три раза. Вписано одно и то же (или на общем листе пусто: файл
    сохранили, не пересчитав формулы) - это одна запись. Разное - какой номер
    верный, KORT не угадывает. Показывается строка листа юрлица: там вписывали.
    """

    def preferred(group: list[Row]) -> Row:
        return next((row for row in group if row.sheet not in SUMMARY_SHEETS), group[0])

    # Дважды на одном листе - тоже не угадывается.
    twins: dict[tuple[str, str], list[Row]] = {}
    for row in rows:
        if row.filled and row.bill is not None:
            twins.setdefault((row.sheet, row.bill.id), []).append(row)
    for (sheet, _), group in twins.items():
        if len(group) > 1:
            where = ", ".join(str(row.row) for row in group)
            for row in group:
                row.fail(f"Этот счёт на листе «{sheet}» несколько раз (строки {where}) - какую брать, KORT не угадывает")

    groups: dict[str, list[Row]] = {}
    order: list[Row | str] = []
    for row in rows:
        if row.bill is None:
            order.append(row)
            continue
        if row.bill.id not in groups:
            groups[row.bill.id] = []
            order.append(row.bill.id)
        groups[row.bill.id].append(row)
    out: list[Row] = []
    for item in order:
        if isinstance(item, Row):
            out.append(item)
            continue
        group = groups[item]
        filled = [row for row in group if row.filled]
        if not filled:
            out.append(preferred(group))
            continue
        failed = [row for row in filled if row.status == ERROR]
        if failed:
            out.append(preferred(failed))
            continue
        chosen = preferred(filled)
        if len({row.entered for row in filled}) > 1:
            where = ", ".join(f"«{row.sheet}» строка {row.row}" for row in filled)
            chosen.fail(f"На листах {where} у этого счёта вписано разное - какой номер и дата верные, KORT не угадывает")
        out.append(chosen)
    return out


def parse(data: bytes, board: Board) -> dict[str, Any]:
    """Строки заполненного шаблона - с тем, что делать с каждой. В книги не пишет.

    `rows` - счета файла по порядку: счёт с нескольких листов - одной строкой
    (`_merge`), пустые строки листа пропущены. `counts` - сколько их в каждом
    состоянии (`STATUSES`).
    """
    book = _open(data)
    try:
        found: list[tuple[str, int, dict[str, int], list[tuple[Any, ...]]]] = []
        reasons: list[tuple[str, str]] = []
        for sheet in book.worksheets:
            try:
                header_row, columns, lines = _find(sheet)
            except TemplateError as exc:
                reasons.append((sheet.title, str(exc)))
                continue
            found.append((sheet.title, header_row, columns, lines))
        if not found:
            if len(reasons) == 1:
                raise TemplateError(reasons[0][1])
            raise TemplateError("; ".join(f"«{title}»: {reason}" for title, reason in reasons))
        everything = [*board.bills, *board.done]
        by_id = {bill.id: bill for bill in everything}
        rows: list[Row] = []
        for title, header_row, columns, lines in found:
            for offset, cells in enumerate(lines):
                # Строка листа: шапка - `header_row + 1`, данные - сразу под ней.
                row = _read_row(title, header_row + 2 + offset, columns, cells, everything, by_id)
                if row is not None:
                    rows.append(row)
    finally:
        book.close()
    if not rows:
        raise TemplateError("В шаблоне нет ни одной строки счёта")
    rows = _merge(rows)
    counts = {status: sum(1 for row in rows if row.status == status) for status in STATUSES}
    return {"rows": [row.out() for row in rows], "counts": counts}


__all__ = ["READY", "SAME", "EMPTY", "ERROR", "TemplateError", "build", "filename", "parse", "period_text"]
