"""Выгрузка реестра в .xlsx — теми же листами, блоками и шапками, что в файле.

Страховка на переходный период и для тех, кому нужен файл. Лист выгрузки —
это отбор реестра: договор стоит в первом подходящем блоке листа, как стоял бы
в Excel. Шапка блока — подписи колонок из файла (у «Заказчик ГК / Заказчик ГК»
заказчик снова в F), текст соглашений — как был, даты — датами Excel, суммы —
числами. Цвета шапки исходного файла возвращаются в выгрузке: на экране их нет,
а в файле человек их ждёт.

Оформление (30.09.2026). До него файл выглядел сломанным: строку с
многострочным текстом (соглашения, примечания) Excel растягивал на 5–10 строк,
остальное прижималось к низу, ссылки Битрикса налезали на соседние колонки,
формат «# ##0» показывал 1 200 000 как «1200 000», а «Оплачено/Остаток по
выписке» выходили пустыми. Теперь: строки одной высоты, многострочное видно
первой строкой (целиком — в ячейке), длинное обрезается краем ячейки, ширины —
по содержимому, суммы — «#,##0», ссылки кликабельные, рамки, фильтр, печать
на ширину листа с шапкой на каждой странице.
"""
from __future__ import annotations

import io
import re
import uuid
from copy import copy
from datetime import date
from decimal import Decimal
from typing import Any, Sequence
from urllib.parse import quote

import sqlalchemy as sa
from sqlalchemy.orm import Session

from app.finance import history
from app.finance.contracts import payments as payments_module
from app.finance.contracts import shares as shares_module
from app.finance.contracts import views as views_module
from app.finance.contracts.fields import CHOICES, LIVE_FIELDS, with_live_columns
from app.finance.contracts.models import Contract, EntityView
from app.finance.contracts.service import (
    Access,
    Actor,
    Registry,
    age_months,
    facts_of,
    people_of,
    visible_to,
)
from app.finance.models import Workspace

ROW_NUMBER = "row_number"
#: Колонки главного листа, если у листа своих колонок нет (реестр заведён
#: не из файла).
DEFAULT_KEYS = (
    ROW_NUMBER, "status", "folder_url", "planned_end_at", "people", "executor", "customer",
    "number", "signed_at", "department", "type", "subject", "amount", "paid_snapshot",
    "remaining_snapshot", "amendments_text", "amendments_summary_text", "end_date", "note",
)

#: Высота строки договора и строки текста в пунктах (Calibri 11 — 15 pt).
ROW_HEIGHT = 18.0
LINE_HEIGHT = 15.0
#: Шапка растёт по подписи, но не больше четырёх строк: длинная подпись из
#: файла («Текущее состояние (действующий/ недействующий/ …)») целиком — в ячейке.
HEADER_LINES = 4
HEAD_FILL = "F2F2F2"
HEAD_LINE = "A6A6A6"
CELL_LINE = "D9D9D9"
LINK_COLOR = "0563C1"
DATE_FORMAT = "DD.MM.YYYY"
DATE_KEYS = ("signed_at", "planned_end_at", "end_date")
MONEY_KEYS = (
    "amount", "paid_snapshot", "remaining_snapshot", "summary_paid", "summary_remaining", *LIVE_FIELDS,
)
#: «Доли исполнителей» — последней колонкой каждого блока, как в листе
#: (`sheet-adapter.ts`). В реестре такого поля нет: доли живут в договоре
#: по людям, и видны — как на экране — только открытые этому человеку.
SHARES_KEY = "__shares"
SHARES_LABEL = "Доли исполнителей"
#: Ширина колонки без ширины из файла — по смыслу; по содержимому она
#: вырастет до потолка, но не сузится.
KIND_WIDTH = {"index": 6.0, "date": 12.0, "money": 14.0, "link": 18.0, "text": 16.0, "shares": 30.0}
KEY_WIDTH = {"customer": 30.0, "executor": 18.0, "subject": 26.0, "number": 16.0, "people": 18.0}
#: Потолок ширины «по содержимому»: дальше текст обрезается краем ячейки.
WIDTH_CAP = {"index": 8.0, "date": 12.0, "money": 18.0, "link": 22.0, "text": 42.0, "shares": 60.0}
_LINK = re.compile(r"^https?://\S+$")


def _text(registry: Registry, contract: Contract, key: str, people: Sequence[uuid.UUID]) -> Any:
    """Значение ячейки выгрузки: так, как его пишут в реестре."""
    if key in ("executor", "customer"):
        party_id = contract.executor_id if key == "executor" else contract.customer_id
        party = registry.parties.get(party_id) if party_id else None
        return party.name if party else None
    if key in ("type", "subject", "status", "economic_role"):
        value_id = getattr(contract, f"{key}_id")
        value = registry.values.get(value_id) if value_id else None
        return value.value if value else None
    if key == "department":
        department = registry.departments.get(contract.department_id) if contract.department_id else None
        return department.code if department else None
    if key == "people":
        names = [registry.employees[item].full_name for item in people if item in registry.employees]
        return ", ".join(names) or None
    if key == "amount":
        if contract.amount is not None:
            return float(contract.amount)
        return contract.amount_terms or None
    if key in ("paid_snapshot", "remaining_snapshot"):
        raw = (contract.file_snapshot or {}).get(key.replace("_snapshot", ""))
        try:
            return float(Decimal(str(raw))) if raw not in (None, "") else None
        except ArithmeticError:
            return raw
    if key in ("summary_paid", "summary_remaining"):
        match = registry.summary_of(contract)
        if match is None or match.state != "found":
            return None
        if key == "summary_paid":
            return float(match.paid)
        return float(contract.amount - match.paid) if contract.amount is not None else None
    if key == "age_months":
        return age_months(contract.signed_at)
    if key in ("signed_at", "planned_end_at", "end_date"):
        return getattr(contract, key)
    if key in ("billing", "end_kind"):
        # Выбор — словами, как на экране: «Расторжение», а не «terminated».
        raw = getattr(contract, key) or None
        return dict(CHOICES[key]).get(raw, raw) if raw else None
    if key in ("number", "folder_url", "amendments_text", "amendments_summary_text", "note", "amount_terms", "currency"):
        return getattr(contract, key) or None
    raw = (contract.attrs or {}).get(key)
    field = registry.field_by_key.get(key)
    if raw in (None, "", []) or field is None:
        return None
    try:
        if field.type in ("list",):
            value = registry.values.get(uuid.UUID(str(raw)))
            return value.value if value else raw
        if field.type == "multi_list":
            return ", ".join(
                registry.values[uuid.UUID(item)].value for item in raw if uuid.UUID(item) in registry.values
            )
        if field.type == "person":
            return ", ".join(
                registry.employees[uuid.UUID(item)].full_name for item in raw if uuid.UUID(item) in registry.employees
            )
        if field.type in ("number", "money"):
            return float(Decimal(str(raw)))
        if field.type == "date":
            return date.fromisoformat(str(raw))
    except (ValueError, KeyError, ArithmeticError):
        return str(raw)
    return raw


def build(session: Session, workspace: Workspace, access: Access, actor: Actor, view_keys: Sequence[str] | None = None) -> bytes:
    """Книга .xlsx: лист реестра — лист книги, блок — название, шапка и строки.

    Книга пишется потоком (`write_only`): строки уходят в файл по мере записи,
    а не копятся объектами ячеек. Обычная книга держала в памяти каждую ячейку
    со стилем — выгрузка 20 000 договоров поднимала процесс на 380 МБ, и
    память обратно не возвращалась. В потоковой книге ширины колонок и
    закрепление шапки задаются до первой строки: их знают заранее из листа.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    registry = Registry(session, workspace)
    contracts = list(
        session.scalars(
            sa.select(Contract)
            .where(Contract.workspace_id == workspace.id, Contract.deleted_at.is_(None))
            .order_by(Contract.position, Contract.created_at)
        )
    )
    people = people_of(session, [item.id for item in contracts])
    visible = [item for item in contracts if visible_to(item, registry, access, people.get(item.id, []))]
    # Стороны всех выгружаемых договоров — одним запросом, а не по одному на
    # договор внутри facts_of.
    registry.parties_for({pid for item in visible for pid in (item.executor_id, item.customer_id)})
    facts = {item.id: facts_of(item, registry, people.get(item.id, [])) for item in visible}
    hidden = set(access.hidden)
    # «Оплачено/Остаток по выписке» — тем же разнесением, что и на экране
    # (`payments.summaries`): по всем договорам компании, показ — по видимым.
    # Кому журнал не открыт, у того этих колонок нет (`access.hidden`).
    live: dict[uuid.UUID, dict[str, Any]] = {}
    if access.view and not hidden.intersection(LIVE_FIELDS):
        allocation = payments_module.allocate(session, workspace, registry, contracts)
        for item in visible:
            summary = payments_module.summary_of(
                item, allocation.paid.get(item.id, []), len(allocation.open.get(item.id, []))
            )
            if summary is not None:
                live[item.id] = summary
    # Доли — тем же ответом, что у экрана: владельцу и администратору — все,
    # начальнику — договоров отдела, сотруднику — только своя.
    shares = shares_module.all_visible(session, workspace, access)["contracts"]

    # Без выбора — листы реестра; «Разовые» выгружаются своей кнопкой.
    views: list[EntityView] = [
        view for view in registry.views if (view.key in view_keys if view_keys else not view.book)
    ]
    book = Workbook(write_only=True)
    taken_titles: list[str] = []
    for view in views:
        title = _sheet_title(view.title, taken_titles)
        taken_titles.append(title)
        sheet = book.create_sheet(title=title)
        styles = _Styles(sheet, (view.style or {}).get("header_fill") or HEAD_FILL)
        blocks = list(view.blocks or [])
        # «Оплачено/Остаток по выписке» — там же, где на экране: за колонками
        # оплат из файла (`with_live_columns`).
        layouts = [
            [
                *(
                    [column for column in with_live_columns(list(block.get("columns") or [])) if column.get("key") not in hidden]
                    or _default_columns(registry, hidden)
                ),
                {"key": SHARES_KEY, "label": SHARES_LABEL},
            ]
            for block in blocks
        ]
        # Договор — в одном блоке листа: в первом подходящем (`views.place`).
        members: list[list[Contract]] = [[] for _ in blocks]
        for item in visible:
            index = views_module.place(view, facts[item.id])
            if index is not None and index < len(members):
                members[index].append(item)
        # Значения — заранее: ширины колонок потоковой книги задаются до
        # первой строки, а считаются по содержимому.
        values: list[list[list[Any]]] = []
        for number in range(len(blocks)):
            rows = []
            for position, item in enumerate(members[number], start=1):
                mine = people.get(item.id, [])
                row: list[Any] = []
                for column in layouts[number]:
                    key = column["key"]
                    if key == ROW_NUMBER:
                        row.append(position)
                    elif key in LIVE_FIELDS:
                        row.append(_number(live.get(item.id, {}).get(key)))
                    elif key == SHARES_KEY:
                        row.append(_shares_text(shares.get(str(item.id)), mine, registry, item.amount))
                    else:
                        row.append(_text(registry, item, key, mine))
                rows.append(row)
            values.append(rows)

        widths: dict[int, float] = {}
        for number, columns in enumerate(layouts):
            for c_index, column in enumerate(columns):
                kind = _kind(registry, column["key"])
                label = column.get("label") or _title(registry, column["key"])
                width = _width(column, kind, label, [row[c_index] for row in values[number]])
                widths[c_index + 1] = max(widths.get(c_index + 1, 0.0), width)
        for c_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(c_index)].width = round(width, 1)
        if len(blocks) == 1:
            header_row = 2 if blocks[0].get("title") else 1
            sheet.freeze_panes = f"A{header_row + 1}"
            last = get_column_letter(max(len(layouts[0]), 1))
            sheet.auto_filter.ref = f"A{header_row}:{last}{header_row + len(values[0])}"
            sheet.print_title_rows = f"{header_row}:{header_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

        line = 0
        for number, block in enumerate(blocks):
            columns = layouts[number]
            kinds = [_kind(registry, column["key"]) for column in columns]
            if number:
                sheet.append([])
                sheet.append([])
                line += 2
            if block.get("title"):
                line += 1
                sheet.row_dimensions[line].height = 22.0
                sheet.append([None, styles.cell(block["title"], "title")])
            labels = [column.get("label") or _title(registry, column["key"]) for column in columns]
            lines = max(
                (
                    _lines(label, widths.get(c_index, KIND_WIDTH["text"]), bold=True)
                    for c_index, label in enumerate(labels, start=1)
                ),
                default=1,
            )
            line += 1
            sheet.row_dimensions[line].height = min(lines, HEADER_LINES) * LINE_HEIGHT + 6
            sheet.append([styles.cell(label, "head") for label in labels])
            for row in values[number]:
                line += 1
                sheet.row_dimensions[line].height = ROW_HEIGHT
                sheet.append([
                    styles.value(value, kinds[c_index], widths.get(c_index + 1, KIND_WIDTH["text"]))
                    for c_index, value in enumerate(row)
                ])
    if not views:
        book.create_sheet("Реестр")
    buffer = io.BytesIO()
    book.save(buffer)
    history.write(
        session,
        workspace,
        kind="contract.export",
        entity="contract",
        title=f"реестр выгружен в .xlsx: {len(visible)} договоров, листов {len(views)}",
        after={"views": [view.key for view in views], "contracts": len(visible)},
        actor=actor.email,
    )
    return buffer.getvalue()


def _default_columns(registry: Registry, hidden: set[str]) -> list[dict[str, Any]]:
    columns = with_live_columns([{"key": key} for key in DEFAULT_KEYS])
    return [
        {"key": column["key"], "label": _title(registry, column["key"])}
        for column in columns
        if column["key"] == ROW_NUMBER or (column["key"] in registry.field_by_key and column["key"] not in hidden)
    ]


def _title(registry: Registry, key: str) -> str:
    if key == ROW_NUMBER:
        return "№"
    field = registry.field_by_key.get(key)
    return field.title if field else key


def _sheet_title(title: str, taken: list[str]) -> str:
    clean = "".join(ch for ch in (title or "Лист") if ch not in "[]:*?/\\")[:31] or "Лист"
    candidate, index = clean, 2
    while candidate in taken:
        suffix = f" {index}"
        candidate = clean[: 31 - len(suffix)] + suffix
        index += 1
    return candidate


def _shares_text(
    entry: dict[str, Any] | None, listed: Sequence[uuid.UUID], registry: Registry, total: Decimal | None
) -> str | None:
    """Доли строкой ячейки — как в листе: «Елжас 500 000 (71,4%) · Рысбек
    200 000 (28,6%)», в конце — что не распределено; сотруднику — «Ваша доля
    …»; совместный договор без сумм — «Елжас · Рысбек — доли не указаны»."""
    if entry is None:
        return None
    people: dict[str, dict[str, str | None]] = entry.get("people") or {}

    def name(employee: str) -> str:
        try:
            found = registry.employees.get(uuid.UUID(employee))
        except ValueError:
            found = None
        return found.full_name if found else "—"

    def one(share: dict[str, str | None] | None) -> str:
        amount = _number((share or {}).get("amount"))
        percent = _number((share or {}).get("percent"))
        pct = f"{percent:.1f}".rstrip("0").rstrip(".").replace(".", ",") + "%" if percent is not None else ""
        if amount is not None:
            return f"{_grouped(amount)} ({pct})" if pct else _grouped(amount)
        return pct or "—"

    order = [str(item) for item in listed]
    if not people:
        if entry.get("scope") == "own":
            return "Ваша доля не указана"
        return f"{' · '.join(name(item) for item in order)} — доли не указаны" if order else None
    if entry.get("scope") == "own":
        return f"Ваша доля {one(next(iter(people.values())))}"
    ids = [*order, *(item for item in people if item not in order)]
    parts = [f"{name(item)} {one(people.get(item)) if item in people else '—'}" for item in ids]
    amounts = [_number(share.get("amount")) for share in people.values()]
    percents = [_number(share.get("percent")) for share in people.values()]
    if total is not None and all(value is not None for value in amounts):
        rest = float(total) - sum(value or 0.0 for value in amounts)
        if rest > 0.5:
            parts.append(f"не распределено {_grouped(round(rest, 2))}")
    elif all(value is not None for value in percents):
        rest = 100 - sum(value or 0.0 for value in percents)
        if rest > 0.01:
            parts.append(f"не распределено {f'{rest:.1f}'.rstrip('0').rstrip('.').replace('.', ',')}%")
    return " · ".join(parts)


def _grouped(value: float) -> str:
    """500000 → «500 000», 1234,5 → «1 234,5» — как в листе."""
    text = f"{value:,.2f}".rstrip("0").rstrip(".")
    return text.replace(",", " ").replace(".", ",")


def _kind(registry: Registry, key: str) -> str:
    """Смысл колонки для оформления: номер строки, дата, сумма, ссылка, текст."""
    if key == ROW_NUMBER:
        return "index"
    if key == SHARES_KEY:
        return "shares"
    if key in DATE_KEYS:
        return "date"
    if key in MONEY_KEYS:
        return "money"
    if key == "folder_url":
        return "link"
    field = registry.field_by_key.get(key)
    if field is not None and field.type == "date":
        return "date"
    if field is not None and field.type in ("number", "money"):
        return "money"
    return "text"


def _number(raw: Any) -> float | None:
    if raw in (None, ""):
        return None
    try:
        return float(Decimal(str(raw)))
    except ArithmeticError:
        return None


def _shown(value: Any) -> str:
    """Как значение будет видно в ячейке — для ширины колонки."""
    if value is None:
        return ""
    if isinstance(value, date):
        return "00.00.0000"
    if isinstance(value, (int, float)):
        return f"{value:,.0f}"
    return str(value).split("\n", 1)[0]


def _width(column: dict[str, Any], kind: str, label: str, values: list[Any]) -> float:
    """Ширина: из файла или по смыслу, дальше — по содержимому до потолка.

    По содержимому — по 90-й доле длин, а не по самой длинной: одна длинная
    ячейка не должна раздувать колонку на всю выгрузку. Подпись шапки не
    рвётся посреди слова: колонка не уже самого длинного слова подписи.
    """
    base = float(column.get("width") or KEY_WIDTH.get(column["key"], KIND_WIDTH[kind]))
    lengths = sorted(len(_shown(value)) for value in values if value not in (None, ""))
    fit = 0.0
    if lengths:
        fit = lengths[min(len(lengths) - 1, int(len(lengths) * 0.9))] * 1.1 + 2
    word = max((len(part) for part in label.split()), default=0) * 1.15 + 2
    width = max(base, min(WIDTH_CAP[kind], max(fit, word)))
    # Подпись шапки — не больше `HEADER_LINES` строк: длинную подпись из
    # файла колонка вмещает, расширяясь (до 34), а не обрезает сверху и снизу.
    while width < 34 and _lines(label, width, bold=True) > HEADER_LINES:
        width += 1
    return width


def _lines(text: str, width: float, *, bold: bool = False) -> int:
    """Сколько строк займёт текст с переносом по словам в колонке этой ширины."""
    capacity = max(1.0, (width - 1) / (1.15 if bold else 1.05))
    total = 0
    for paragraph in str(text).split("\n"):
        count, used = 1, 0.0
        for word in paragraph.split():
            size = float(len(word))
            if used and used + 1 + size > capacity:
                count += 1
                used = 0.0
            while size > capacity:
                count += 1
                size -= capacity
            used = used + (1 if used else 0) + size
        total += count
    return total


class _Styles:
    """Готовые стили ячеек потоковой книги.

    Стиль собирается один раз на сочетание и копируется в ячейку: присваивать
    каждой из сотен тысяч ячеек шрифт, рамку и выравнивание заново — это
    поиск в таблице стилей книги на каждое присваивание.
    """

    def __init__(self, sheet: Any, head_fill: str) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self.sheet = sheet
        thin = Side(style="thin", color=CELL_LINE)
        head = Side(style="thin", color=HEAD_LINE)
        grid = Border(left=thin, right=thin, top=thin, bottom=thin)
        link = Font(color=LINK_COLOR, underline="single")
        self.recipes: dict[str, dict[str, Any]] = {
            "title": {"font": Font(bold=True, size=12), "alignment": Alignment(vertical="center")},
            "head": {
                "font": Font(bold=True),
                "fill": PatternFill("solid", start_color=head_fill),
                "border": Border(left=head, right=head, top=head, bottom=head),
                "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
            },
            "empty": {"border": grid},
            "text": {"border": grid, "alignment": Alignment(vertical="center")},
            # Длиннее колонки: «заполнить» — Excel обрезает текст краем
            # ячейки, а не выводит его поверх соседних пустых.
            "clip": {"border": grid, "alignment": Alignment(horizontal="fill", vertical="center")},
            # Несколько строк: видна первая, остальное — в ячейке. Перенос
            # включён, но высота строки своя — Excel её не растягивает.
            "lines": {"border": grid, "alignment": Alignment(vertical="top", wrap_text=True)},
            "link": {"border": grid, "font": link, "alignment": Alignment(horizontal="fill", vertical="center")},
            "link_short": {"border": grid, "font": link, "alignment": Alignment(vertical="center")},
            "index": {"border": grid, "alignment": Alignment(horizontal="center", vertical="center")},
            "date": {
                "border": grid,
                "number_format": DATE_FORMAT,
                "alignment": Alignment(horizontal="center", vertical="center"),
            },
            "money": {
                "border": grid,
                "number_format": "#,##0",
                "alignment": Alignment(horizontal="right", vertical="center"),
            },
            "cents": {
                "border": grid,
                "number_format": "#,##0.00",
                "alignment": Alignment(horizontal="right", vertical="center"),
            },
        }
        self._made: dict[str, Any] = {}

    def cell(self, value: Any, name: str) -> Any:
        from openpyxl.cell import WriteOnlyCell

        cell = WriteOnlyCell(self.sheet, value=value)
        style = self._made.get(name)
        if style is None:
            for attr, item in self.recipes[name].items():
                setattr(cell, attr, item)
            self._made[name] = copy(cell._style)
        else:
            cell._style = copy(style)
        return cell

    def value(self, value: Any, kind: str, width: float) -> Any:
        if value is None or value == "":
            return self.cell(None, "empty")
        if isinstance(value, date):
            return self.cell(value, "date")
        if isinstance(value, bool):
            return self.cell(value, "text")
        if isinstance(value, (int, float)):
            if kind == "index":
                return self.cell(value, "index")
            return self.cell(value, "cents" if float(value) % 1 else "money")
        text = str(value)
        if "\n" in text:
            return self.cell(text, "lines")
        long = len(text) * 1.05 > width - 1
        if kind == "link" and _LINK.match(text):
            cell = self.cell(text, "link" if long else "link_short")
            cell.hyperlink = _target(text)
            return cell
        return self.cell(text, "clip" if long else "text")


def _target(url: str) -> str:
    """Адрес ссылки для Excel: кириллица и пробелы — percent-кодом, иначе
    Excel при открытии просит «восстановить содержимое»."""
    return quote(url, safe=":/?#[]@!$&'()*+,;=%~-._")


__all__ = ["build"]
