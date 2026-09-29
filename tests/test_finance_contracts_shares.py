"""Доли исполнителей и отделов, изменения листа и возврат «как было» (29.09.2026).

Набор держит то, что было обещано владельцу:

* сотрудник видит только свою долю — чужих сумм нет в ответе вовсе;
  начальник отдела, администратор и владелец видят и правят все;
* людей в договоре сколько угодно, доли не больше суммы договора и 100%;
* правка поля «Ответственное лицо» не стирает доли оставшихся;
* доли отделов — только администратору и начальнику с полным правом;
* «История» договора не приносит чужие суммы;
* ломающее изменение листа ставит точку восстановления, и её возвращает
  автор или администратор — не трогая то, что после поменяли коллеги.
"""
from __future__ import annotations

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from test_finance_access import BASE, activate, client, department, employee, grant, register


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.core.config import settings
    from app.finance import audit, auth
    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    monkeypatch.setattr(settings, "environment", "test")
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()
    yield engine
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()


@pytest.fixture
def app(finance_db) -> FastAPI:
    from finance_routes import finance_app

    return finance_app()


def _team(app: FastAPI):
    """Владелец; ЮО: начальник, два юриста и юрист без входа; НО: человек с правом на все договоры."""
    owner = register(app)
    yuo, no = department(owner, "ЮО"), department(owner, "НО")
    head = employee(owner, "Начальникова Жанель", "+77051120001", yuo)
    first = employee(owner, "Юристов Рысбек", "+77051120002", yuo)
    second = employee(owner, "Юристов Тимур", "+77051120003", yuo)
    other = employee(owner, "Налоговая Дана", "+77051120004", no)
    quiet = owner.post(f"{BASE}/people/employees", json={"full_name": "Салтанат", "department_id": yuo})
    assert quiet.status_code == 201, quiet.text
    grant(owner, "department", yuo, {"contracts": {"level": "edit", "scope": {"rows": "department"}}})
    grant(owner, "department", no, {"contracts": "edit"})
    grant(owner, "employee", head["id"], {"people": {"level": "edit", "scope": {"rows": "department"}}})
    people = {
        "owner": owner,
        "head": activate(app, "+77051120001", ip="10.0.1.1"),
        "first": activate(app, "+77051120002", ip="10.0.1.2"),
        "second": activate(app, "+77051120003", ip="10.0.1.3"),
        "other": activate(app, "+77051120004", ip="10.0.1.4"),
    }
    ids = {"head": head["id"], "first": first["id"], "second": second["id"], "quiet": quiet.json()["id"], "other": other["id"]}
    return people, ids, yuo, no


def _contract(owner: TestClient, number: str, amount: str = "700000", people: str = "Юристов Рысбек, Юристов Тимур, Салтанат", **extra):
    values = {"number": number, "amount": amount, "people": people, "department": "ЮО", **extra}
    response = owner.post(f"{BASE}/contracts", json={"values": values})
    assert response.status_code == 201, response.text
    return response.json()["contract"]["id"]


def _shares(who: TestClient, contract_id: str) -> dict:
    response = who.get(f"{BASE}/contracts/{contract_id}/shares")
    assert response.status_code == 200, response.text
    return response.json()


def _put(who: TestClient, contract_id: str, unit: str, items: list[tuple[str, str | None]], kind: str = "people"):
    key = "employee_id" if kind == "people" else "department_id"
    return who.put(
        f"{BASE}/contracts/{contract_id}/shares/{kind}",
        json={"unit": unit, "items": [{key: item, "value": value} for item, value in items]},
    )


def test_dolyu_vidit_kazhdyy_svoyu_nachalnik_i_vladelets_vse(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/700")
    done = _put(owner, contract, "amount", [(ids["first"], "500 000"), (ids["second"], "150000"), (ids["quiet"], "50000")])
    assert done.status_code == 200, done.text
    full = done.json()["people"]
    assert full["scope"] == "all" and full["can_edit"] is True and full["unit"] == "amount"
    assert [row["amount"] for row in full["rows"]] == ["500000.00", "150000.00", "50000.00"]
    assert full["summary"]["allocated_amount"] == "700000.00" and full["summary"]["rest_amount"] == "0.00"
    assert full["summary"]["over"] is False

    # Юрист — только своя доля: чужих строк в ответе нет вовсе.
    mine = _shares(team["first"], contract)["people"]
    assert mine["scope"] == "own" and mine["can_edit"] is False and mine["count"] == 3
    assert len(mine["rows"]) == 1 and mine["rows"][0]["mine"] is True
    assert mine["rows"][0]["amount"] == "500000.00" and mine["rows"][0]["percent"].startswith("71.42")
    assert "summary" not in mine
    assert _put(team["first"], contract, "amount", [(ids["first"], "700000")]).status_code == 403

    # Начальник отдела — все доли своего отдела, и правит их.
    boss = _shares(team["head"], contract)["people"]
    assert boss["scope"] == "all" and boss["can_edit"] is True and len(boss["rows"]) == 3
    assert _put(team["head"], contract, "percent", [(ids["first"], "60"), (ids["second"], "40")]).status_code == 200

    # Человек другого отдела с правом на все договоры — ни одной доли.
    stranger = _shares(team["other"], contract)
    assert stranger["people"]["scope"] == "none" and stranger["people"]["rows"] == []
    assert stranger["departments"] is None

    # Проценты следуют за суммой договора.
    after = _shares(owner, contract)["people"]
    assert after["unit"] == "percent"
    assert [row["amount"] for row in after["rows"]] == ["420000.00", "280000.00", None]


def test_doli_ne_bolshe_summy_i_sta_protsentov(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/701")
    over_amount = _put(owner, contract, "amount", [(ids["first"], "500000"), (ids["second"], "300000")])
    assert over_amount.status_code == 400 and "больше суммы договора" in over_amount.json()["detail"]
    over_percent = _put(owner, contract, "percent", [(ids["first"], "60"), (ids["second"], "50")])
    assert over_percent.status_code == 400 and "больше 100%" in over_percent.json()["detail"]
    negative = _put(owner, contract, "amount", [(ids["first"], "-5")])
    assert negative.status_code == 400
    stranger = _put(owner, contract, "amount", [(ids["other"], "5")])
    assert stranger.status_code == 400 and "ответственного" in stranger.json()["detail"]
    # Не всё распределено — это не ошибка, а остаток.
    part = _put(owner, contract, "amount", [(ids["first"], "500000")])
    assert part.status_code == 200
    assert part.json()["people"]["summary"]["rest_amount"] == "200000.00"


def test_pravka_otvetstvennyh_ne_stiraet_doli(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/702")
    assert _put(owner, contract, "amount", [(ids["first"], "400000"), (ids["second"], "300000")]).status_code == 200
    # Третьего убрали, первого и второго поменяли местами — доли остались у людей.
    patched = owner.patch(
        f"{BASE}/contracts/{contract}", json={"values": {"people": "Юристов Тимур, Юристов Рысбек"}}
    )
    assert patched.status_code == 200, patched.text
    rows = {row["employee_id"]: row["amount"] for row in _shares(owner, contract)["people"]["rows"]}
    assert rows == {ids["second"]: "300000.00", ids["first"]: "400000.00"}
    # Ушедший из поля уходит со своей долей.
    owner.patch(f"{BASE}/contracts/{contract}", json={"values": {"people": "Юристов Тимур"}})
    rows = {row["employee_id"]: row["amount"] for row in _shares(owner, contract)["people"]["rows"]}
    assert rows == {ids["second"]: "300000.00"}


def test_istoriya_ne_nesyot_chuzhie_summy(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/703")
    assert _put(owner, contract, "amount", [(ids["first"], "500000"), (ids["second"], "200000")]).status_code == 200
    kinds = lambda who: [item["kind"] for item in who.get(f"{BASE}/contracts/{contract}/history").json()["items"]]  # noqa: E731
    assert "contract.shares.people" in kinds(owner)
    assert "contract.shares.people" in kinds(team["head"])
    assert "contract.shares.people" not in kinds(team["first"])
    # Заголовок события — без сумм: его видят в журнале действий.
    item = next(i for i in owner.get(f"{BASE}/contracts/{contract}/history").json()["items"] if i["kind"] == "contract.shares.people")
    assert "500" not in item["title"] and "задано у 2 из 3" in item["title"]


def test_svodka_dolej_otdaet_tolko_otkrytoe(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/704")
    _contract(owner, "ЮО/705", people="Юристов Тимур")
    _put(owner, contract, "amount", [(ids["first"], "500000"), (ids["second"], "200000")])
    everything = owner.get(f"{BASE}/contracts/shares").json()["contracts"]
    assert set(everything) == {contract} and len(everything[contract]["people"]) == 2
    mine = team["first"].get(f"{BASE}/contracts/shares").json()["contracts"]
    assert mine[contract]["scope"] == "own" and list(mine[contract]["people"]) == [ids["first"]]
    assert team["other"].get(f"{BASE}/contracts/shares").json()["contracts"] == {}


def test_sovmestnyy_dogovor_bez_summ_v_otbore_s_dolyami(app: FastAPI) -> None:
    """30.09: на проде долей не ввели ни в одном договоре, и «С долями» у
    владельца был пуст, хотя совместных договоров шесть. Совместный — в
    ответе и без сумм; каждому — в своих пределах."""
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    together = _contract(owner, "ЮО/720", people="Юристов Рысбек, Юристов Тимур")
    alone = _contract(owner, "ЮО/721", people="Юристов Тимур")
    everything = owner.get(f"{BASE}/contracts/shares").json()["contracts"]
    assert everything == {together: {"scope": "all", "people": {}}}
    head = team["head"].get(f"{BASE}/contracts/shares").json()["contracts"]
    assert head == {together: {"scope": "all", "people": {}}}
    # Исполнитель — свой договор, чужих сумм нет; посторонний — ничего.
    first = team["first"].get(f"{BASE}/contracts/shares").json()["contracts"]
    assert first == {together: {"scope": "own", "people": {}}}
    assert team["other"].get(f"{BASE}/contracts/shares").json()["contracts"] == {}
    # Доля у одного — у второго исполнителя по-прежнему только своя.
    _put(owner, together, "amount", [(ids["first"], "400000")])
    second = team["second"].get(f"{BASE}/contracts/shares").json()["contracts"]
    assert second == {together: {"scope": "own", "people": {}}}
    first = team["first"].get(f"{BASE}/contracts/shares").json()["contracts"]
    assert list(first[together]["people"]) == [ids["first"]]
    assert alone not in owner.get(f"{BASE}/contracts/shares").json()["contracts"]


def test_doli_v_vygruzke_excel_po_pravam(app: FastAPI) -> None:
    """30.09: в скачанном .xlsx колонки «Доли исполнителей» не было вовсе.
    Теперь — последней, как в листе, и видно в ней то же, что на экране."""
    import io

    from openpyxl import load_workbook

    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    split = _contract(owner, "ЮО/730", people="Юристов Рысбек, Юристов Тимур")
    _contract(owner, "ЮО/731", people="Юристов Рысбек, Юристов Тимур")
    _put(owner, split, "amount", [(ids["first"], "500000"), (ids["second"], "150000")])

    def column(who: TestClient) -> dict[str, str | None]:
        response = who.get(f"{BASE}/contracts/export.xlsx")
        assert response.status_code == 200, response.text
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        head = rows[0]
        assert head[-1] == "Доли исполнителей"
        number = head.index("№ Договора")
        return {str(row[number]): row[-1] for row in rows[1:]}

    everything = column(owner)
    assert everything["ЮО/730"] == "Юристов Рысбек 500 000 (71,4%) · Юристов Тимур 150 000 (21,4%) · не распределено 50 000"
    assert everything["ЮО/731"] == "Юристов Рысбек · Юристов Тимур — доли не указаны"
    first = column(team["first"])
    assert first["ЮО/730"] == "Ваша доля 500 000 (71,4%)" and first["ЮО/731"] == "Ваша доля не указана"
    assert "150 000" not in str(first)
    assert team["other"].get(f"{BASE}/contracts/export.xlsx").status_code == 200


def test_doli_otdelov_tolko_nachalniku_s_polnym_pravom(app: FastAPI) -> None:
    team, ids, yuo, no = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/706")
    done = _put(owner, contract, "percent", [(yuo, "70"), (no, "30")], kind="departments")
    assert done.status_code == 200, done.text
    rows = done.json()["departments"]["rows"]
    assert [(row["code"], row["amount"], row["main"]) for row in rows] == [("ЮО", "490000.00", True), ("НО", "210000.00", False)]
    # Начальник ЮО с договорами «своего отдела» — не «включено всё»: долей отделов нет.
    assert _shares(team["head"], contract)["departments"] is None
    assert _put(team["head"], contract, "percent", [(yuo, "100")], kind="departments").status_code == 403
    # Дали всё — открылись.
    grant(owner, "employee", ids["head"], {"contracts": {"level": "edit", "scope": {"rows": "all", "beyond": True}}})
    assert _shares(team["head"], contract)["departments"]["rows"][0]["code"] == "ЮО"
    # Не начальник с тем же правом — нет.
    grant(owner, "employee", ids["first"], {"contracts": {"level": "edit", "scope": {"rows": "all", "beyond": True}}})
    assert _shares(team["first"], contract)["departments"] is None
    # Отдел, названный дважды, и больше 100% — отказ.
    twice = _put(owner, contract, "percent", [(yuo, "50"), (yuo, "20")], kind="departments")
    assert twice.status_code == 400
    # Пустой список — отделов в договоре нет.
    cleared = _put(owner, contract, "amount", [], kind="departments")
    assert cleared.status_code == 200 and cleared.json()["departments"]["rows"] == []


# ── Изменения листа и точки восстановления ──────────────────────────────────


def _change(who: TestClient, **body):
    return who.post(f"{BASE}/contracts/sheet/change", json=body)


def test_udalenie_iz_lista_vozvrashchaet_avtor(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    one, two = _contract(owner, "ЮО/801"), _contract(owner, "ЮО/802")
    lawyer = team["first"]
    done = _change(lawyer, action="delete_contracts", ids=[one, two], view="main")
    assert done.status_code == 200, done.text
    assert sorted(done.json()["done"]) == sorted([one, two]) and done.json()["point"]["kind"] == "contracts_delete"
    point = done.json()["point"]["id"]
    assert owner.get(f"{BASE}/contracts/{one}").status_code == 404
    # Точку видит автор и администратор, коллега — нет.
    assert [item["id"] for item in lawyer.get(f"{BASE}/contracts/restore-points").json()["items"]] == [point]
    assert team["second"].get(f"{BASE}/contracts/restore-points").json()["items"] == []
    assert team["second"].post(f"{BASE}/contracts/restore-points/{point}/restore").status_code == 403
    back = lawyer.post(f"{BASE}/contracts/restore-points/{point}/restore")
    assert back.status_code == 200, back.text
    assert len(back.json()["done"]) == 2
    assert owner.get(f"{BASE}/contracts/{one}").status_code == 200
    again = lawyer.post(f"{BASE}/contracts/restore-points/{point}/restore")
    assert again.status_code == 400 and "уже возвращено" in again.json()["detail"]
    listed = owner.get(f"{BASE}/contracts/restore-points").json()["items"][0]
    assert listed["restored_at"] and listed["can_restore"] is False


def test_vozvrat_znacheniy_ne_trogaet_chuzhuyu_pravku(app: FastAPI) -> None:
    team, ids, _yuo, _no = _team(app)
    owner = team["owner"]
    one, two = _contract(owner, "ЮО/811", note="было-1"), _contract(owner, "ЮО/812", note="было-2")
    point = _change(
        owner, action="values_point", items=[{"id": one, "keys": ["note", "number"]}, {"id": two, "keys": ["note"]}]
    ).json()["point"]["id"]
    owner.patch(f"{BASE}/contracts/{one}", json={"values": {"note": "вставка", "number": "ЮО/811-б"}})
    owner.patch(f"{BASE}/contracts/{two}", json={"values": {"note": "вставка"}})
    # После вставки коллега поправил примечание второго договора.
    team["first"].patch(f"{BASE}/contracts/{two}", json={"values": {"note": "коллега"}})
    back = owner.post(f"{BASE}/contracts/restore-points/{point}/restore")
    assert back.status_code == 200, back.text
    first_values = owner.get(f"{BASE}/contracts/{one}").json()["contract"]["values"]
    assert first_values["note"] == "было-1" and first_values["number"] == "ЮО/811"
    assert owner.get(f"{BASE}/contracts/{two}").json()["contract"]["values"]["note"] == "коллега"
    assert any("поменяли" in text for text in back.json()["skipped"])


def test_kolonka_iz_lista_i_vozvrat(app: FastAPI) -> None:
    team, _ids, _yuo, _no = _team(app)
    owner = team["owner"]
    _contract(owner, "ЮО/821")
    added = _change(owner, action="add_column", view="main", title="Источник клиента", after=["number"])
    assert added.status_code == 200, added.text
    key = added.json()["key"]
    schema = owner.get(f"{BASE}/contracts/schema").json()
    main = next(view for view in schema["views"] if view["key"] == "main")
    keys = [column["key"] for column in main["blocks"][0]["columns"]]
    assert keys[0] == "row_number" and keys[keys.index("number") + 1] == key
    # Юристу колонки не открыты.
    assert _change(team["first"], action="add_column", view="main", title="Своё").status_code == 403
    # Переименовали шапку, убрали колонку — и вернули по точкам в обратном порядке.
    renamed = _change(owner, action="rename_column", view="main", block=0, key="number", label="Номер")
    assert renamed.status_code == 200, renamed.text
    removed = _change(owner, action="remove_column", view="main", keys=["note"])
    assert removed.status_code == 200 and removed.json()["removed"]
    for point in (removed.json()["point"]["id"], renamed.json()["point"]["id"], added.json()["point"]["id"]):
        assert owner.post(f"{BASE}/contracts/restore-points/{point}/restore").status_code == 200
    schema = owner.get(f"{BASE}/contracts/schema").json()
    main = next(view for view in schema["views"] if view["key"] == "main")
    assert main["blocks"][0]["columns"] == []
    assert key not in {field["key"] for field in schema["fields"]}
    # «№» не убирается.
    assert _change(owner, action="remove_column", view="main", keys=["row_number"]).status_code == 400


def test_list_pereimenovan_i_ubran_vozvrat(app: FastAPI) -> None:
    team, _ids, _yuo, _no = _team(app)
    owner = team["owner"]
    view = owner.post(
        f"{BASE}/contracts/setup/views", json={"title": "Аренда", "blocks": [{"title": "", "filter": {"any": []}}]}
    ).json()
    renamed = _change(owner, action="rename_view", view=view["key"], title="Аренда ГК")
    assert renamed.status_code == 200, renamed.text
    archived = _change(owner, action="archive_view", view=view["key"])
    assert archived.status_code == 200
    assert _change(owner, action="archive_view", view="main").status_code == 400
    assert owner.post(f"{BASE}/contracts/restore-points/{archived.json()['point']['id']}/restore").status_code == 200
    assert owner.post(f"{BASE}/contracts/restore-points/{renamed.json()['point']['id']}/restore").status_code == 200
    titles = {item["key"]: item["title"] for item in owner.get(f"{BASE}/contracts/schema").json()["views"]}
    assert titles[view["key"]] == "Аренда"


def test_stroki_peredvinuty_i_vernulis(app: FastAPI) -> None:
    team, _ids, _yuo, _no = _team(app)
    owner = team["owner"]
    first, second, third = (_contract(owner, f"ЮО/83{i}") for i in range(3))
    order = lambda: [item["id"] for item in sorted(owner.get(f"{BASE}/contracts").json()["contracts"], key=lambda c: c["position"])]  # noqa: E731
    assert order() == [first, second, third]
    moved = _change(team["first"], action="move_rows", ids=[third], before=first, view="main")
    assert moved.status_code == 200, moved.text
    assert order() == [third, first, second]
    assert owner.post(f"{BASE}/contracts/restore-points/{moved.json()['point']['id']}/restore").status_code == 200
    assert order() == [first, second, third]


def test_vstavlennaya_stroka_vstaet_pered_sosedom(app: FastAPI) -> None:
    team, _ids, _yuo, _no = _team(app)
    owner = team["owner"]
    first, second = _contract(owner, "ЮО/841"), _contract(owner, "ЮО/842")
    made = owner.post(f"{BASE}/contracts", json={"values": {"number": "ЮО/840"}, "before": second})
    assert made.status_code == 201, made.text
    listed = sorted(owner.get(f"{BASE}/contracts").json()["contracts"], key=lambda item: item["position"])
    assert [item["values"]["number"] for item in listed] == ["ЮО/841", "ЮО/840", "ЮО/842"]
    assert listed[0]["id"] == first and listed[2]["id"] == second


def test_istoriya_pishet_vybor_slovami_a_ne_klyuchom(app: FastAPI) -> None:
    """30.09: в журнал уходило «смысл даты окончания: — → terminated»."""
    team, _, _, _ = _team(app)
    owner = team["owner"]
    contract = _contract(owner, "ЮО/830")
    response = owner.patch(f"{BASE}/contracts/{contract}", json={"values": {"end_kind": "terminated"}})
    assert response.status_code == 200, response.text
    titles = [item["title"] for item in owner.get(f"{BASE}/contracts/{contract}/history").json()["items"]]
    assert any("— → Расторжение" in title for title in titles), titles
    assert not any("terminated" in title for title in titles), titles
