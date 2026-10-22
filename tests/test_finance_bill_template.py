"""Шаблон счёта Excel (`contracts/bill_template.py`): скачать, заполнить, загрузить.

Набор держит обещания шаблона:
* колонки и порядок - как в шаблоне финансистов BBC (22.10.2026), известное
  вписано, «№ Счета» и «Дата выстав. Счета» - жёлтые и пустые;
* загруженный шаблон разбирается по названиям колонок, строка счёта - по
  скрытому коду KORT или по юрлицу, клиенту и № договора; шаблон финансистов
  на несколько счетов тоже подходит;
* не нашлась строка, нашлось две, нет номера или даты, правили договор, не
  .xlsx - отказ словами, а не догадка;
* пустые «Счет за Период» и «Доп. Коммент.» не стирают книгу;
* скачать и загрузить можно только с правом «Счета: правка»; скачивание - в журнале.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.finance.contracts import bill_template
from app.finance.contracts.bills import Bill

AYA = Bill(
    line=10, group="overdue", text="Просрочка 2–5 дней", own="BBCS", customer="ИП АЙЯ", number="№НО/39",
    month="ОКТЯБРЬ 2026", month_no=10, kind="Абон.П.", staff="Ануар", department="НО",
    amount=Decimal("91000"), period_from=date(2026, 10, 20), period_to=date(2026, 11, 19), paid="", status="",
    note="", onec_name="ИП Айя", signed_at="20.01.2026",
)


def workbook(data: bytes):
    return load_workbook(io.BytesIO(data))


def filled(data: bytes, number=312, issued=datetime(2026, 10, 22)) -> bytes:
    book = workbook(data)
    sheet = book.active
    sheet["N2"].value = number
    sheet["O2"].value = issued
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_shablon_kak_u_finansistov() -> None:
    sheet = workbook(bill_template.build(AYA)).active
    heads = {cell.column_letter: cell.value for cell in sheet[1] if cell.value}
    assert heads == {
        "B": "Наша\nФирма", "C": "Наш\nСотрудник", "D": "Заказчик\n(Название Фирмы)", "E": "По 1C",
        "F": "№\nДоговора", "G": "Дата заключения \nДоговора", "I": "Период\nс (нач.)", "J": "Период\nпо (зав.)",
        "L": "Сумма\nДоговора", "N": "№ \nСчета", "O": "Дата выстав.\nСчета", "Q": "Код KORT",
    }
    row = {cell.column_letter: cell.value for cell in sheet[2] if cell.value is not None}
    assert row["B"] == "BBCS" and row["C"] == "Ануар" and row["D"] == "ИП АЙЯ" and row["E"] == "ИП Айя"
    assert row["F"] == "№НО/39" and row["L"] == 91000
    assert row["G"].date() == date(2026, 1, 20) and row["I"].date() == date(2026, 10, 20)
    assert row["J"].date() == date(2026, 11, 19)
    assert "N" not in row and "O" not in row, "номер и дату заполняет человек"
    assert sheet["N1"].fill.fgColor.rgb == "FFFFF2CC" and sheet["O2"].fill.fgColor.rgb == "FFFFF2CC", "жёлтое, не прозрачное"
    assert sheet["O2"].number_format == "DD.MM.YYYY"
    assert row["Q"] == AYA.id and sheet.column_dimensions["Q"].hidden, "код счёта - в скрытой колонке"
    assert bill_template.filename(AYA) == "Счёт ИП АЙЯ №НО-39 ОКТЯБРЬ 2026.xlsx"


def test_zapolnennyy_shablon_v_polya_kartochki() -> None:
    result = bill_template.parse(filled(bill_template.build(AYA)), AYA)
    assert result == {"values": {"bill_no": "312", "bill_date": "2026-10-22"}, "sheet": "Счёт", "row": 2}
    # Номер текстом и дата текстом - тоже.
    text = filled(bill_template.build(AYA), number="248/1", issued="22.10.2026")
    assert bill_template.parse(text, AYA)["values"] == {"bill_no": "248/1", "bill_date": "2026-10-22"}


def finance_template(rows: list[list]) -> bytes:
    """Шаблон в точности как прислали финансисты 22.10.2026: шапка в 1-й строке, без кода KORT."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист1"
    heads = {
        "B": "Наша\nФирма", "C": "Наш\nСотрудник", "D": "Заказчик\n(Название Фирмы)", "E": "По 1C",
        "F": "№\nДоговора", "G": "Дата заключения \nДоговора", "I": "Период\nс (нач.)", "J": "Период\nпо (зав.)",
        "L": "Сумма\nДоговора", "N": "№ \nСчета", "O": "Дата выстав.\nСчета",
    }
    for letter, title in heads.items():
        sheet[f"{letter}1"] = title
    for number, values in enumerate(rows, start=2):
        for letter, value in zip(("B", "C", "D", "E", "F", "G", "I", "J", "L", "N", "O"), values):
            sheet[f"{letter}{number}"] = value
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


THREE = [
    ["BBCS", "Дана", "ИП Кадыркулова Торгын Карашевна", "ИП Кадыркулова", "№НО/19", datetime(2026, 1, 17),
     datetime(2026, 10, 17), datetime(2026, 11, 16), "25\xa0000", 501, datetime(2026, 10, 22)],
    ["BBCS", "Ануар", "ИП АЙЯ", "ИП Айя", "№НО/39", datetime(2026, 1, 20),
     datetime(2026, 10, 20), datetime(2026, 11, 19), "91\xa0000", 502, datetime(2026, 10, 22)],
    ["BBCS", "Ануар", "ИП РАЖА", "ИП Ража", "№НО/37", datetime(2026, 1, 20),
     datetime(2026, 10, 20), datetime(2026, 11, 19), "78\xa0000", None, None],
]


def test_shablon_finansistov_na_neskolko_schetov() -> None:
    result = bill_template.parse(finance_template(THREE), AYA)
    assert result["values"] == {"bill_no": "502", "bill_date": "2026-10-22"} and result["row"] == 3
    raja = Bill(**{**AYA.__dict__, "customer": "ИП РАЖА", "number": "№НО/37"})
    with pytest.raises(bill_template.TemplateError, match="не вписан «№ Счета»"):
        bill_template.parse(finance_template(THREE), raja)
    stranger = Bill(**{**AYA.__dict__, "customer": "ТОО KENESARY AGRO", "number": "№ЮО/44", "own": "BBCL"})
    with pytest.raises(bill_template.TemplateError, match="нет строки этого счёта: BBCL, ТОО KENESARY AGRO"):
        bill_template.parse(finance_template(THREE), stranger)
    # Тот же клиент за прошлый месяц - другая строка: по «Периоду с» не путается.
    september = Bill(**{**AYA.__dict__, "period_from": date(2026, 9, 20), "month": "СЕНТЯБРЬ 2026"})
    with pytest.raises(bill_template.TemplateError, match="нет строки этого счёта"):
        bill_template.parse(finance_template(THREE), september)


def test_dva_raza_otkaz() -> None:
    twice = finance_template([THREE[1], THREE[1]])
    with pytest.raises(bill_template.TemplateError, match="несколько раз"):
        bill_template.parse(twice, AYA)


def test_pravili_dogovor_v_shablone() -> None:
    book = workbook(filled(bill_template.build(AYA)))
    book.active["D2"].value = "ИП Другой"
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(bill_template.TemplateError, match="другой договор, чем в карточке"):
        bill_template.parse(out.getvalue(), AYA)


def test_bez_daty_i_plokhaya_data() -> None:
    with pytest.raises(bill_template.TemplateError, match="пусто"):
        bill_template.parse(filled(bill_template.build(AYA), issued=None), AYA)
    with pytest.raises(bill_template.TemplateError, match="не дата: «завтра»"):
        bill_template.parse(filled(bill_template.build(AYA), issued="завтра"), AYA)


def test_ne_xlsx_otkaz() -> None:
    with pytest.raises(bill_template.TemplateError, match="не файл .xlsx"):
        bill_template.parse(b"%PDF-1.4 not a workbook", AYA)
    with pytest.raises(bill_template.TemplateError, match="пустой"):
        bill_template.parse(b"", AYA)
    with pytest.raises(bill_template.TemplateError, match="больше 2 МБ"):
        bill_template.parse(b"PK" + b"0" * (bill_template.MAX_BYTES + 1), AYA)


def test_kolonki_po_nazvaniyam_i_dop_polya() -> None:
    """Вставили колонку и добавили «Счет за Период» и «Доп. Коммент.»: читается по названиям; пустое не стирает."""
    book = workbook(filled(bill_template.build(AYA)))
    sheet = book.active
    sheet.insert_cols(1)
    sheet["S1"], sheet["S2"] = "Счет за\nПериод (Коммент)", "с 20.10.2026 года по 19.11.2026 года"
    sheet["T1"], sheet["T2"] = "Доп. \nКоммент.", None
    out = io.BytesIO()
    book.save(out)
    values = bill_template.parse(out.getvalue(), AYA)["values"]
    assert values == {"bill_no": "312", "bill_date": "2026-10-22", "bill_period": "с 20.10.2026 года по 19.11.2026 года"}
