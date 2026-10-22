"""«Счета» и бот счетов: книга-сводка → KORT и канал Telegram, карточка счёта → книга юрлица.

Google и Telegram подменены: проверяется разбор книги и что бот шлёт, а не
сеть. Набор держит обещания раздела и бота:

* колонка «Счет» находится по названию рядом с «№ Счета», «Счет за Период» и
  «Счет2» и после вставленной колонки; состояние в KORT - словами книги,
  незнакомое - в «Без срока», а не угадано;
* два поста - абонентская плата и аренда, не вперемешку; без служебных строк
  шаблона; день счёта по «Периоду с»; «Разовый» и «Прочее» не идут;
  вчерашнее и просроченное - с пометкой; без вида и без даты - названы
  строкой; имя клиента - ссылка на его карточку в KORT; длинный пост режется
  с самых старых;
* посты по часам: второй проход и второй процесс не выкладывают их снова,
  отказ Telegram освобождает час, отметка правит пост, всё выставлено - пост
  удалён;
* новое - сразу и со звуком, в пост своего вида, но не чаще `NEW_GAP`;
  кнопка «Прислать сейчас» - изменившиеся посты без звука и не чаще `PULL_GAP`,
  итог на 5 секунд в канале, старые нажатия не отрабатываются;
  `/bind` ставит кнопку заново;
* переход с одного поста на два правит висящий пост и выкладывает пост аренды
  без звука, а не выдаёт весь список новым;
* карточка счёта пишет в книгу юрлица, список и канал видят запись сразу,
  запись - в журнале; без права «Счета: правка» записать нельзя;
* раздел открыт только с правом «Счета», без подключённой сводки - пусто.
"""
from __future__ import annotations

import time
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from urllib.parse import unquote

import pytest
import sqlalchemy as sa
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.finance import auth, bills_bot
from app.finance.config import finance_settings
from app.finance.contracts import bill_book, bills
from app.finance.contracts.models import BillNotice
from app.finance.db import finance_session
from app.finance.support_bot import TelegramError

BASE = "/api/v1/finance"
PASSWORD = "pass-12345"
CHAT = "@bbc_bills"

HEADER = [
    "Техн. 2Наша \nФирма",
    "Мес",
    "Техн. 1",
    "Наша \nФирма",
    "Вид \nУслуги",
    "Наш \nСотрудник",
    "Заказчик\n(Название Фирмы)",
    "Сумма\nДоговора",
    "Отдел",
    "№\nДоговора",
    "Период\nс (нач.)",
    "Период\nпо (зав.)",
    "Доп. \nКоммент.",
    "Счет",
    "Счет за\nПериод (Коммент)",
    "№ \nСчета",
    "Дата выстав.\nСчета",
    "Факт Оплата",
    "Сумма Факт\nПоступ.",
    "Статус\n(Продл./Приостн./Разв./На исп.)",
    "Счет2",
]
BILL = HEADER.index("Счет")
BILL_NO = HEADER.index("№ \nСчета")
BILL_DATE = HEADER.index("Дата выстав.\nСчета")


def until_of(period: str) -> str:
    """«22.10.2026» → «21.11.2026»: месяц договора, как в книге."""
    if not period:
        return ""
    start = datetime.strptime(period, "%d.%m.%Y").date()
    year, month = (start.year + 1, 1) if start.month == 12 else (start.year, start.month + 1)
    return (date(year, month, start.day) - timedelta(days=1)).strftime("%d.%m.%Y")


def row(
    month_no: int,
    month: str,
    own: str,
    customer: str,
    number: str,
    bill: str,
    *,
    kind: str = "Абон.П.",
    department: str = "ОБО",
    amount: str = "100 000",
    period: str = "",
    staff: str = "Меруерт",
    paid: str = "FALSE",
    bill_no: str = "",
    bill_date: str = "",
) -> list[str]:
    return [
        "", str(month_no), month, own, kind, staff, customer, amount, department, number,
        period, until_of(period), "", bill, "", bill_no, bill_date, paid, "0", "", bill,
    ]


def rent(month_no: int, month: str, customer: str, number: str, bill: str, **extra: str) -> list[str]:
    extra.setdefault("kind", "Аренда")
    extra.setdefault("department", "Аренда")
    extra.setdefault("amount", "50 000")
    return row(month_no, month, "BBC Omar", customer, number, bill, staff='ТОО "BBC Omar"', **extra)


OCT = "ОКТЯБРЬ 2026"
SHEET = [
    HEADER,
    rent(6, "ИЮНЬ 2026", "ТОО Молфуд", "№RENT/93", "Просрочка более 5 дней", period="04.06.2026"),
    rent(7, "ИЮЛЬ 2026", "ТОО Молфуд", "№RENT/93", "Просрочка более 5 дней", period="04.07.2026"),
    rent(8, "АВГУСТ 2026", "ТОО Молфуд", "№RENT/93", "Просрочка более 5 дней", period="04.08.2026"),
    rent(10, OCT, "ТОО Молфуд", "№RENT/93", "Просрочка более 5 дней", period="04.10.2026"),
    row(10, OCT, "BBC", "ADK Home", "№ОБО/110", "Просрочка 2–5 дней", kind="нет", period="17.10.2026", paid="Часть"),
    row(10, OCT, "BBC", "Алатау Пауэр", "№ОБО/146", "TRUE", period="25.10.2026", bill_no="502", bill_date="21.10.2026"),
    row(10, OCT, "BBCL", "TrioSynergy", "№ЮО/58", "3 дня до выставления", department="ЮО", period="24.10.2026"),
    row(10, OCT, "BBCL", "Meliora", "№ЮО/35-2", "Еще рано", department="ЮО", period="26.10.2026"),
    row(8, "АВГУСТ 2026", "BBCL", "TrioSynergy", "№ЮО/136", "", kind="нет", department="ЮО"),
    row(5, "Старые (до Май/Июня 2026)", "BBCS", "Интер Табыс", "№НО/108", "FALSE", kind="нет", department="НО",
        period="15.05.2026"),
    row(10, OCT, "BBCL", "ТОО KENESARY AGRO", "№ЮО/44", "Сегодня день выставления", department="ЮО",
        amount="150 000", period="22.10.2026"),
    rent(10, OCT, "ТОО Custom Trade", "№RENT/84", "Просрочка 1 день", kind="нет", amount="0", period="21.10.2026"),
    row(10, OCT, "BBCS", "ИП АЙЯ", "№НО/39", "Просрочка 2–5 дней", department="НО", amount="91 000",
        period="20.10.2026"),
    row(10, OCT, "EA", "ТОРИ КОМПАНИ", "№НО/123", "Просрочка более 5 дней", kind="Разовый", department="НО",
        amount="500 000", period="15.10.2026"),
    rent(10, OCT, "ИП ЖАЛМҰХАН", "№ RENT/112", ""),
    rent(10, OCT, "ТОО Find Dream", "№RENT/94", "Сегодня день выставления", amount="", period="22.10.2026"),
    ["", "", "", "", "", "", "", "", "", "", "", "", "", "FALSE"],  # пустая строка листа - не счёт
]


def line_of(values: list[list[str]], customer: str, period: str = "") -> list[str]:
    """Строка книги по клиенту (и «Периоду с») - тест правит её, как человек правит книгу."""
    start = HEADER.index("Период\nс (нач.)")
    found = [line for line in values[1:] if customer in line and (not period or line[start] == period)]
    assert len(found) == 1, (customer, period, len(found))
    return found[0]


def link(own: str, customer: str, number: str, month: str = OCT) -> str:
    """Имя клиента в посте - ссылкой на карточку счёта в KORT."""
    bill_id = bills.bill_id(bills.bill_key(own, customer, number, month))
    return f'<a href="https://kort.cards/?s=invoices&amp;bill={bill_id}">{customer}</a>'


def test_sostoyanie_slovami_knigi() -> None:
    assert bills.state_of("TRUE") is None and bills.state_of("ИСТИНА") is None
    assert bills.state_of("Просрочка 1 день") == bills.OVERDUE
    assert bills.state_of("1 день до выставления") == bills.SOON
    assert bills.state_of("Сегодня день выставления") == bills.SOON, "день счёта - не «без срока»"
    assert bills.state_of("Ещё рано") == bills.EARLY
    for strange in ("", "FALSE", "Не нужен"):
        assert bills.state_of(strange) == bills.UNDATED, strange


def test_kniga_razbiraetsya_po_nazvaniyam() -> None:
    board = bills.build_board(SHEET, title="Осн.Общая сводка BBC 2026", worksheet="Сводка все ЮР лица")
    assert board.counts() == {bills.OVERDUE: 8, bills.SOON: 3, bills.EARLY: 1, bills.UNDATED: 3}
    assert board.rows == 16, "пустая строка листа - не договор"
    assert all(bill.customer != "Алатау Пауэр" for bill in board.bills), "выставленный счёт не в списке"
    [done] = board.done
    assert (done.customer, done.bill_no, done.bill_date, done.group) == ("Алатау Пауэр", "502", "21.10.2026", bills.DONE)
    first = board.bills[0]
    assert (first.customer, first.month, first.line) == ("ТОО Молфуд", "ИЮНЬ 2026", 2), "старое - первым"
    adk = next(bill for bill in board.bills if bill.customer == "ADK Home")
    assert adk.text == "Просрочка 2–5 дней" and adk.paid == "part" and adk.amount == Decimal("100000")
    assert adk.period_from == date(2026, 10, 17) and adk.period_to == date(2026, 11, 16)
    assert board.total(bills.OVERDUE) == Decimal("891000")
    assert {bill.text for bill in board.of(bills.UNDATED)} == {"", "FALSE"}
    assert board.find(adk.id) is adk and board.find(done.id) is done and len(adk.id) == 12

    # В книгу вставили колонку перед «Счет»: читается то же самое.
    shifted = [line[:5] + ["вставили"] + line[5:] for line in SHEET]
    again = bills.build_board(shifted, title="", worksheet="Сводка все ЮР лица")
    assert [bill.out()["text"] for bill in again.bills] == [bill.out()["text"] for bill in board.bills]
    assert "«Счет»" in again.drift


def test_bez_kolonok_scheta_list_chitaetsya() -> None:
    """«Счет за Период» находилась «похожей» на «Счет», и лист без неё не читался (22.10.2026)."""
    short = [
        ["Наша Фирма", "Заказчик (Название Фирмы)", "№ Договора", "Техн. 1", "Счет"],
        ["BBCL", "ТОО KENESARY AGRO", "№ЮО/44", OCT, "Сегодня день выставления"],
    ]
    board = bills.build_board(short, title="", worksheet="Сводка")
    [bill] = board.bills
    assert (bill.customer, bill.bill_no, bill.bill_period) == ("ТОО KENESARY AGRO", "", "")


def test_bez_kolonki_schet_otkaz() -> None:
    broken = [[cell for cell in HEADER if cell != "Счет"]]
    with pytest.raises(bills.BillsError, match="Счет"):
        bills.build_board(broken, title="", worksheet="Сводка")


def test_zapisannoe_iz_kort_vidno_srazu() -> None:
    """Сводка догоняет книгу с задержкой: до тех пор список верит записи KORT."""
    workspace_id = uuid.uuid4()
    board = bills.build_board(SHEET, title="", worksheet="Сводка")
    kenesary = next(bill for bill in board.bills if bill.customer == "ТОО KENESARY AGRO")
    bills.mark(workspace_id, kenesary.key, {"bill_no": "312", "bill_date": "22.10.2026", "bill_period": "", "note": ""})
    seen = bills.view(workspace_id, board)
    assert kenesary.key in seen.issued and all(bill.key != kenesary.key for bill in seen.bills)
    assert seen.find(kenesary.id).bill_no == "312" and seen.find(kenesary.id).group == bills.DONE
    # Сводка, прочитанная после записи, показала то же - отметка больше не нужна.
    later = [list(line) for line in SHEET]
    line_of(later, "ТОО KENESARY AGRO")[BILL] = "TRUE"
    fresh_board = bills.build_board(later, title="", worksheet="Сводка")
    assert bills.view(workspace_id, fresh_board) is fresh_board
    # Снятая отметка возвращает счёт в список по дате, пока сводка не перечитана.
    alatau = next(bill for bill in board.done if bill.customer == "Алатау Пауэр")
    bills.mark(workspace_id, alatau.key, {"bill_no": "", "bill_date": "", "bill_period": "", "note": ""})
    back = bills.view(workspace_id, board).find(alatau.id)
    assert back.group in bills.GROUPS and back.bill_no == ""
    bills.forget(workspace_id)


NOW = datetime(2026, 10, 22, 5, 30, tzinfo=timezone.utc)  # четверг, 10:30 в Астане


def board_of(values: list[list[str]]) -> bills.Board:
    return bills.build_board(values, title="Осн.Общая сводка BBC 2026", worksheet="Сводка")


def test_vid_uslugi() -> None:
    board = board_of(SHEET)
    kinds = {bill.customer: bills_bot.kind_of(bill) for bill in board.bills}
    assert kinds["ТОО KENESARY AGRO"] == bills_bot.ABON
    assert kinds["ТОО Find Dream"] == bills_bot.RENT
    assert kinds["ТОО Custom Trade"] == bills_bot.RENT, "вида нет - аренду узнаём по отделу"
    assert kinds["ТОРИ КОМПАНИ"] == bills_bot.SKIP, "разовые не присылаются"
    assert kinds["ADK Home"] == bills_bot.UNKNOWN, "вида нет и отдел не аренда - не угадываем"


ABON_POST = "\n".join(
    [
        "<b>Абонентская плата: счета к выставлению (22.10.2026)</b>",
        "",
        "<b>BBCL</b> · 22.10 - 21.11",
        f"1) {link('BBCL', 'ТОО KENESARY AGRO', '№ЮО/44')} - №ЮО/44 = 150 000 ₸",
        "",
        "<b>BBCS</b> · просрочено",
        f"1) {link('BBCS', 'ИП АЙЯ', '№НО/39')} - №НО/39 = 91 000 ₸ · с 20.10",
        "",
        "<b>ВСЕГО: 2 счёта = 241 000 ₸</b>",
        "",
        "Без вида услуги в книге: ADK Home",
    ]
)

RENT_POST = "\n".join(
    [
        "<b>Аренда: счета к выставлению (22.10.2026)</b>",
        "",
        "<b>BBC Omar</b> · 22.10 - 21.11",
        f"1) {link('BBC Omar', 'ТОО Find Dream', '№RENT/94')} - №RENT/94",
        "",
        "<b>BBC Omar</b> · вчерашнее, 21.10 - 20.11",
        f"1) {link('BBC Omar', 'ТОО Custom Trade', '№RENT/84')} - №RENT/84 = 0 ₸",
        "",
        "<b>BBC Omar</b> · просрочено",
        f"1) {link('BBC Omar', 'ТОО Молфуд', '№RENT/93')} - №RENT/93 = 50 000 ₸ · с 04.10",
        "",
        "<b>ВСЕГО: 3 счёта = 50 000 ₸</b>",
        "",
        "Раньше сентября: ещё 3 счёта на 150 000 ₸ - список в KORT",
        "Без даты «Период с» в книге: ИП ЖАЛМҰХАН",
    ]
)


def test_dva_posta_bez_sluzhebnyh_strok(monkeypatch) -> None:
    """Абонентская плата и аренда - отдельными постами; ни черт, ни «шаблон v1.3»."""
    monkeypatch.setattr(finance_settings, "app_url", "https://kort.cards")
    board = board_of(SHEET)
    abon = bills_bot.render(board, NOW, bills_bot.ABON)
    rent_post = bills_bot.render(board, NOW, bills_bot.RENT)
    assert abon is not None and rent_post is not None
    assert abon.text == ABON_POST
    assert rent_post.text == RENT_POST
    for text in (abon.text, rent_post.text):
        assert "---" not in text and "шаблон" not in text
        assert chr(0x2014) not in text, "длинного тире в посте нет (CLAUDE.md)"
    assert "TrioSynergy" not in abon.text and "Meliora" not in abon.text, "день счёта ещё не наступил"
    assert "ТОРИ" not in abon.text + rent_post.text, "разовые не присылаются"
    assert "Алатау" not in abon.text, "выставлен"


def test_dlinnyy_post_rezhetsya_so_starogo(monkeypatch) -> None:
    monkeypatch.setattr(finance_settings, "app_url", "https://kort.cards")
    many = [HEADER] + [
        rent(10, OCT, f"ТОО Арендатор с длинным названием {index}", f"№RENT/{index}", "Просрочка",
             period=f"{index % 20 + 1:02d}.10.2026")
        for index in range(150)
    ] + [line_of(SHEET, "ТОО Find Dream")]
    post = bills_bot.render(board_of(many), NOW, bills_bot.RENT)
    assert post is not None
    assert bills_bot._visible(post.text) <= 4096 and bills_bot._entities(post.text) <= bills_bot.ENTITIES
    assert "ТОО Find Dream" in post.text, "сегодняшнее не режется"
    assert "<b>ВСЕГО: 151 счёт = 7 500 000 ₸</b>" in post.text, "итог - по всем, и по не вошедшим"
    assert "· с 20.10" in post.text and "· с 01.10" not in post.text, "уходят самые старые"
    hidden = int(post.text.split("и ещё ", 1)[1].split(" ", 1)[0])
    assert post.text.count("ТОО Арендатор с длинным названием") + hidden == 150


def test_pustoy_spisok_ne_vykladyvaetsya() -> None:
    nothing = board_of([HEADER, line_of(SHEET, "Алатау Пауэр"), line_of(SHEET, "Meliora")])
    assert bills_bot.render(nothing, NOW, bills_bot.ABON) is None
    assert bills_bot.render(nothing, NOW, bills_bot.RENT) is None
    tail = board_of([HEADER, line_of(SHEET, "ТОО Молфуд", "04.06.2026")])
    assert bills_bot.render(tail, NOW, bills_bot.RENT) is None, "одно старое - только первым постом дня и по кнопке"
    alone = bills_bot.render(tail, NOW, bills_bot.RENT, extras_alone=True)
    assert alone is not None and "Раньше сентября: ещё 1 счёт на 50 000 ₸" in alone.text
    assert "ВСЕГО" not in alone.text
    assert bills_bot.render(tail, NOW, bills_bot.ABON, extras_alone=True) is None, "аренда в пост абонентской платы не идёт"


# ── Посты по часам ───────────────────────────────────────────────────────────


class FakeBot:
    def __init__(self) -> None:
        self.calls: list[tuple[str, dict]] = []
        self.fail: TelegramError | None = None
        self.refuse: dict[str, TelegramError] = {}
        self.next_id = 100

    def call(self, method: str, params: dict | None = None, **_: object) -> dict:
        if self.fail is not None:
            raise self.fail
        if method in self.refuse:
            raise self.refuse[method]
        self.calls.append((method, params or {}))
        if method == "sendMessage":
            self.next_id += 1  # номер - только у нового сообщения, как в Telegram
        return {"message_id": self.next_id}

    def sent(self, method: str) -> list[dict]:
        return [params for name, params in self.calls if name == method]

    def posts(self, title: str) -> list[dict]:
        return [item for item in self.sent("sendMessage") if item["text"].startswith(f"<b>{title}")]


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.core.config import settings
    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    monkeypatch.setattr(settings, "environment", "test")
    monkeypatch.setattr(finance_settings, "support_bot_token", "")
    monkeypatch.setattr(finance_settings, "bills_bot_token", "test-token")
    monkeypatch.setattr(finance_settings, "bills_chat_id", CHAT)
    monkeypatch.setattr(finance_settings, "bills_workspace_id", "")
    monkeypatch.setattr(finance_settings, "bills_hour", "10")
    monkeypatch.setattr(finance_settings, "bills_every", "3")
    monkeypatch.setattr(finance_settings, "app_url", "https://kort.cards")
    auth._ATTEMPTS.clear()
    bills._slots.clear()
    bills._marks.clear()
    bill_book.forget()
    yield engine
    bills._slots.clear()
    bills._marks.clear()
    bill_book.forget()
    auth._ATTEMPTS.clear()


@pytest.fixture
def app(finance_db) -> FastAPI:
    from finance_routes import finance_app

    return finance_app()


def client(app: FastAPI) -> TestClient:
    return TestClient(app, headers={"x-client-ip": "10.0.0.7", "user-agent": "pytest-agent"})


def register(app: FastAPI) -> TestClient:
    person = client(app)
    response = person.post(
        f"{BASE}/auth/register",
        json={"email": "owner@bbc.kz", "password": PASSWORD, "company": "BBC (тест)", "full_name": "Ермеков Нурболат"},
    )
    assert response.status_code == 201, response.text
    return person


@pytest.fixture
def book(monkeypatch):
    """Книга в памяти: тест правит `book["values"]` - «в книге поставили галочку»."""
    from app.finance.contracts import summary

    state = {"values": [list(line) for line in SHEET]}
    monkeypatch.setattr(
        summary, "_read",
        lambda spreadsheet_id, tab: summary.build_index(state["values"], title="Осн.Общая сводка BBC 2026", worksheet=tab),
    )
    monkeypatch.setattr(
        bills, "_read",
        lambda spreadsheet_id, tab: bills.build_board(
            state["values"], title="Осн.Общая сводка BBC 2026", worksheet=tab, spreadsheet_id=spreadsheet_id
        ),
    )
    return state


def connect(owner: TestClient) -> None:
    response = owner.post(
        f"{BASE}/contracts/setup/summary",
        json={"link": "https://docs.google.com/spreadsheets/d/abc123/edit", "worksheet": "Сводка все ЮР лица"},
    )
    assert response.status_code == 200, response.text


def fresh() -> None:
    """Книгу перечитать сразу: бот берёт прочитанную, свежая доедет к следующему проходу."""
    with bills._guard:
        slots = list(bills._slots.values())
    for slot in slots:
        board = bills._read(*slot.source)
        with bills._guard:
            slot.board, slot.at = board, time.monotonic()


def issue_all(values: list[list[str]]) -> None:
    for line in values[1:]:
        if len(line) > BILL:
            line[BILL] = "TRUE"


BOTH_POSTED = "abon-posted+rent-posted"
BOTH_SAME = "abon-same+rent-same"


def test_posty_po_chasu_i_pravka(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()

    assert bills_bot.tick(api, CHAT, NOW) == BOTH_POSTED
    abon, rent_post = api.sent("sendMessage")
    assert abon["text"] == ABON_POST and rent_post["text"] == RENT_POST, "абонентская плата - первой"
    assert abon["chat_id"] == CHAT and abon["parse_mode"] == "HTML"
    assert abon["disable_notification"] is None, "по часу - со звуком"
    assert abon["reply_markup"]["inline_keyboard"] == [
        [
            {"text": "Прислать сейчас", "callback_data": "bills:now"},
            {"text": "Открыть в KORT", "url": "https://kort.cards/?s=invoices"},
        ]
    ]

    # Второй проход и «второй процесс» (тот же час в базе) - не выкладывают снова.
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=1)) == BOTH_SAME
    for kind in bills_bot.KINDS:
        assert bills_bot.claim(bills_bot.target_workspace(), bills_bot.slot_of(NOW), kind) is None
    assert len(api.sent("sendMessage")) == 2 and not api.sent("editMessageText")

    # Отметили аренду - правится только пост аренды.
    line_of(book["values"], "ТОО Молфуд", "04.10.2026")[BILL] = "TRUE"
    fresh()
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=5)) == "abon-same+rent-edited"
    [edit] = api.sent("editMessageText")
    assert edit["message_id"] == 102 and "ТОО Молфуд" not in edit["text"]
    assert "<b>ВСЕГО: 2 счёта = 0 ₸</b>" in edit["text"]
    fresh()
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=6)) == BOTH_SAME

    # Выставили всё - постов в канале больше нет.
    issue_all(book["values"])
    fresh()
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=7)) == "abon-deleted+rent-deleted"
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 101}, {"chat_id": CHAT, "message_id": 102}]
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=8)) == "wait"

    # Завтра - новые посты, и вчерашнее в них - с пометкой.
    book["values"] = [list(line) for line in SHEET]
    fresh()
    assert bills_bot.tick(api, CHAT, NOW + timedelta(days=1)) == BOTH_POSTED
    tomorrow = api.posts("Абонентская плата")[-1]["text"]
    assert tomorrow.startswith("<b>Абонентская плата: счета к выставлению (23.10.2026)</b>")
    assert "<b>BBCL</b> · вчерашнее, 22.10 - 21.11" in tomorrow
    assert f"{link('BBC Omar', 'ТОО Custom Trade', '№RENT/84')} - №RENT/84 = 0 ₸ · с 21.10" in (
        api.posts("Аренда")[-1]["text"]
    )


def test_otkaz_telegram_osvobozhdaet_chas(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    api.fail = TelegramError("нет связи с Telegram")
    with pytest.raises(TelegramError):
        bills_bot.tick(api, CHAT, NOW)
    api.fail = None
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=1)) == BOTH_POSTED


def test_posty_kazhdyy_den_v_svoy_chas(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    early = datetime(2026, 10, 22, 3, 0, tzinfo=timezone.utc)  # 08:00 в Астане
    late = datetime(2026, 10, 22, 16, 0, tzinfo=timezone.utc)  # 21:00
    saturday = datetime(2026, 10, 24, 5, 0, tzinfo=timezone.utc)  # 10:00
    # Первый проход запоминает список молча: он придёт постами по часу, а не «новым».
    assert bills_bot.tick(api, CHAT, early) == "seeded"
    assert bills_bot.tick(api, CHAT, late) == "wait", "по часу после 20:00 не выходит"
    assert api.calls == []
    assert bills_bot.tick(api, CHAT, saturday) == BOTH_POSTED, "и в выходные"
    assert len(api.sent("sendMessage")) == 2


def test_napominanie_kazhdye_tri_chasa(app: FastAPI, book) -> None:
    """10, 13, 16, 19: новые посты звенят, прошлые удаляются; строка - если что-то изменилось."""
    owner = register(app)
    connect(owner)
    api = FakeBot()
    morning = datetime(2026, 10, 22, 5, 0, tzinfo=timezone.utc)  # 10:00 в Астане
    assert bills_bot.tick(api, CHAT, morning) == BOTH_POSTED
    assert bills_bot.tick(api, CHAT, morning + timedelta(hours=2, minutes=59)) == BOTH_SAME, "12:59 - ещё 10:00"

    assert bills_bot.tick(api, CHAT, morning + timedelta(hours=3)) == BOTH_POSTED
    second = api.posts("Абонентская плата")[-1]
    assert second["text"].startswith("<b>Абонентская плата: счета к выставлению (22.10.2026)</b>\n\n"), (
        "ничего не изменилось - строки нет"
    )
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 101}, {"chat_id": CHAT, "message_id": 102}]

    line_of(book["values"], "ИП АЙЯ")[BILL] = "TRUE"
    book["values"].append(
        row(10, OCT, "BBC Omar", "ИП АЛТЫНБЕКОВА", "№ОБО/92", "Сегодня день выставления", amount="800 000",
            period="22.10.2026")
    )
    fresh()
    # Пост по часу назвал АЛТЫНБЕКОВУ - отдельного «нового» о ней нет.
    assert bills_bot.tick(api, CHAT, morning + timedelta(hours=6)) == BOTH_POSTED
    third = api.posts("Абонентская плата")[-1]["text"]
    assert "\n<i>С 13:00 выставлено 1, добавилось 1</i>\n" in third
    assert "<i>" not in api.posts("Аренда")[-1]["text"], "в аренде не менялось"
    assert bills_bot.tick(api, CHAT, morning + timedelta(hours=6, minutes=1)) == BOTH_SAME

    late = morning + timedelta(hours=10)  # 20:00 - по часам больше нет
    posts = len(api.sent("sendMessage"))
    assert "posted" not in bills_bot.tick(api, CHAT, late)
    assert "posted" not in bills_bot.tick(api, CHAT, morning + timedelta(hours=14))
    assert len(api.sent("sendMessage")) == posts
    assert bills_bot.slot_of(morning + timedelta(hours=24)) == 2026102310
    assert bills_bot.tick(api, CHAT, morning + timedelta(hours=24)) == BOTH_POSTED


NIGHT = datetime(2026, 10, 21, 22, 0, tzinfo=timezone.utc)  # 03:00 в Астане, четверг 22.10


def test_novoe_srazu_v_svoy_post(app: FastAPI, book, monkeypatch) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    assert bills_bot.tick(api, CHAT, NIGHT) == "seeded"

    # В книгу добавили счёт абонентской платы, чей день сегодня, - его пост сразу, ночью и со звуком.
    book["values"].append(
        row(10, OCT, "BBCS", "ИП НУСУПБАЕВ", "№ОБО/100", "Сегодня день выставления", amount="150 000",
            period="22.10.2026")
    )
    fresh()
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=1)) == "abon-new"
    [first] = api.sent("sendMessage")
    assert first["text"].startswith("<b>Абонентская плата")
    assert f"<b>BBCS</b> · 22.10 - 21.11\n1) {link('BBCS', 'ИП НУСУПБАЕВ', '№ОБО/100')} - №ОБО/100 = 150 000 ₸" in first["text"]
    assert "ТОО KENESARY AGRO" in first["text"], "в посте - весь список вида, а не одно новое"
    assert "Custom Trade" not in first["text"], "аренда - в своём посте"
    assert first["disable_notification"] is None, "ночью - со звуком (решение владельца)"
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=2)) == "abon-same"

    # Новое в аренде - пост аренды, абонентский не трогается.
    book["values"].append(rent(10, OCT, "ТОО YI DA", "№RENT/100", "Сегодня день выставления", period="22.10.2026"))
    fresh()
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=2, seconds=30)) == "rent-new+abon-same"
    assert api.sent("sendMessage")[-1]["text"].startswith("<b>Аренда")

    # Ещё новое через минуту - правка, а не второй звонок.
    book["values"].append(
        row(10, OCT, "BBC Omar", "ТОО Green Line Customs", "№ОБО/56", "Сегодня день выставления", amount="50 000",
            period="22.10.2026")
    )
    fresh()
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=3)) == "abon-edited+rent-same"
    assert len(api.sent("sendMessage")) == 2 and "Green Line Customs" in api.sent("editMessageText")[-1]["text"]

    # Выставили - строка уходит; сняли отметку после паузы - счёт снова новый.
    monkeypatch.setattr(bills_bot, "NEW_GAP", 0.0)
    kenesary = line_of(book["values"], "ТОО KENESARY AGRO")
    kenesary[BILL] = "TRUE"
    fresh()
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=4)) == "abon-edited+rent-same"
    assert "KENESARY" not in api.sent("editMessageText")[-1]["text"]
    kenesary[BILL] = "Сегодня день выставления"
    fresh()
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=5)) == "abon-new+rent-same"
    assert "KENESARY" in api.sent("sendMessage")[-1]["text"]
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 101}], "прошлый пост вида - прочь"

    # Полночь: наступил день следующих счетов - пост сразу, вчерашнее - с пометкой.
    book["values"].append(
        row(10, OCT, "BBCS", "ИП ПЕН Г.И.", "№НО/47", "1 день до выставления", department="НО", amount="34 000",
            period="23.10.2026")
    )
    fresh()
    midnight = datetime(2026, 10, 22, 19, 0, 30, tzinfo=timezone.utc)  # 00:00:30 в Астане, 23.10
    assert "new" not in bills_bot.tick(api, CHAT, midnight - timedelta(minutes=1))
    assert bills_bot.tick(api, CHAT, midnight).startswith("abon-new")
    night_post = api.posts("Абонентская плата")[-1]["text"]
    assert night_post.startswith("<b>Абонентская плата: счета к выставлению (23.10.2026)</b>")
    assert f"<b>BBCS</b> · 23.10 - 22.11\n1) {link('BBCS', 'ИП ПЕН Г.И.', '№НО/47')}" in night_post
    assert "<b>BBCL</b> · вчерашнее, 22.10 - 21.11" in night_post


def test_novoe_ne_teryaetsya_pri_otkaze_telegram(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    assert bills_bot.tick(api, CHAT, NIGHT) == "seeded"
    book["values"].append(
        row(10, OCT, "BBCS", "ИП НУСУПБАЕВ", "№ОБО/100", "Сегодня день выставления", period="22.10.2026")
    )
    fresh()
    api.fail = TelegramError("нет связи с Telegram")
    with pytest.raises(TelegramError):
        bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=1))
    api.fail = None
    assert bills_bot.tick(api, CHAT, NIGHT + timedelta(minutes=2)) == "abon-new"
    assert "ИП НУСУПБАЕВ" in api.sent("sendMessage")[0]["text"]


def test_staryy_post_ne_udalit_pravitsya(app: FastAPI, book) -> None:
    """Telegram не удаляет сообщения старше 48 часов - тогда «Всё выставлено.» текстом."""
    owner = register(app)
    connect(owner)
    api = FakeBot()
    assert bills_bot.tick(api, CHAT, NOW) == BOTH_POSTED
    issue_all(book["values"])
    fresh()
    api.refuse["deleteMessage"] = TelegramError("Bad Request: message can't be deleted", code=400)
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=2)) == "abon-edited+rent-edited"
    assert api.sent("editMessageText")[-1]["text"].endswith("\n\nВсё выставлено.")
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=3)) == BOTH_SAME


def test_perehod_s_odnogo_posta_na_dva(app: FastAPI, book) -> None:
    """На проде висит общий пост (22.10, 16:00): он становится постом абонентской платы, аренда выходит рядом без звука."""
    owner = register(app)
    connect(owner)
    workspace_id = bills_bot.target_workspace()
    bills_bot._ensure(workspace_id)
    for name, value in (("bills.day", 2026102216), ("bills.post", 8), ("bills.hash", 12345), ("bills.seeded", 2)):
        bills_bot._set(workspace_id, name, value)
    with finance_session() as session:
        session.add(BillNotice(workspace_id=workspace_id, key="старая-запись", bill_group="due", line=""))
    api = FakeBot()
    afternoon = datetime(2026, 10, 22, 11, 30, tzinfo=timezone.utc)  # 16:30 в Астане
    assert bills_bot.tick(api, CHAT, afternoon) == "seeded+rent-posted+abon-edited"
    [rent_post] = api.sent("sendMessage")
    assert rent_post["text"] == RENT_POST and rent_post["disable_notification"] is True, "без звука"
    [edit] = api.sent("editMessageText")
    assert edit["message_id"] == 8 and edit["text"] == ABON_POST
    assert bills_bot._value(workspace_id, "bills.day.rent") == 2026102216, "второй пост аренды в 16:00 не выйдет"
    assert bills_bot.tick(api, CHAT, afternoon + timedelta(minutes=1)) == BOTH_SAME
    with finance_session() as session:
        assert session.get(BillNotice, (workspace_id, "старая-запись", "due")) is not None, "сказанное - помнится"


def test_pervyy_zapusk_posle_soobshcheniy_novoe(app: FastAPI, book) -> None:
    """Записи сообщений «новое» (soon/overdue) стираются, нового поста нет."""
    owner = register(app)
    connect(owner)
    workspace_id = bills_bot.target_workspace()
    bills_bot._ensure(workspace_id)
    bills_bot._set(workspace_id, "bills.seeded", 1)
    with finance_session() as session:
        session.add(BillNotice(workspace_id=workspace_id, key="x", bill_group="soon", line=""))
    api = FakeBot()
    assert bills_bot.tick(api, CHAT, NIGHT) == "seeded"
    with finance_session() as session:
        groups = set(session.scalars(sa.select(BillNotice.bill_group)))
    assert groups == {"due"} and api.calls == []


# ── Кнопка «Прислать сейчас» и /bind ─────────────────────────────────────────


COUNTS = "абонентская плата - 2 счёта, аренда - 3 счёта"


def test_knopka_prislat_seychas(app: FastAPI, book, monkeypatch) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    assert bills_bot.pull(api, CHAT, NOW) == f"Списки обновлены: {COUNTS}."
    abon, rent_post = api.sent("sendMessage")
    assert abon["text"] == ABON_POST and rent_post["text"] == RENT_POST
    assert abon["disable_notification"] is True and rent_post["disable_notification"] is True, "по кнопке - без звука"
    assert bills_bot.pull(api, CHAT, NOW) == "Списки только что обновлены - они выше."
    assert len(api.sent("sendMessage")) == 2

    # Книга не менялась - посты не перевыкладываются, только проверяется, что они на месте.
    monkeypatch.setattr(bills_bot, "PULL_GAP", 0.0)
    assert bills_bot.pull(api, CHAT, NOW + timedelta(minutes=1)) == f"Новых счетов нет, списки выше актуальны: {COUNTS}."
    assert len(api.sent("sendMessage")) == 2 and api.sent("deleteMessage") == []
    assert [item["message_id"] for item in api.sent("editMessageText")] == [101, 102]

    # Пост удалили из канала руками - кнопка выкладывает его заново.
    api.refuse["editMessageText"] = TelegramError("Bad Request: message to edit not found", code=400)
    assert bills_bot.pull(api, CHAT, NOW + timedelta(minutes=1)) == f"Списки обновлены: {COUNTS}."
    assert [item["text"] for item in api.sent("sendMessage")[2:]] == [ABON_POST, RENT_POST]
    assert [item["message_id"] for item in api.sent("deleteMessage")] == [101, 102]
    del api.refuse["editMessageText"]

    # Пост по кнопке не делает час: в 10:00 выйдет свой.
    assert bills_bot.tick(api, CHAT, NOW + timedelta(minutes=2)) == BOTH_POSTED

    issue_all(book["values"])
    fresh()
    assert bills_bot.pull(api, CHAT, NOW + timedelta(minutes=3)) == "На сегодня все счета выставлены."
    assert [item["message_id"] for item in api.sent("deleteMessage")][-2:] == [105, 106]


def test_knopka_posle_pravki_knigi(app: FastAPI, book, monkeypatch) -> None:
    """Книга изменилась с прошлого поста - кнопка выкладывает списки заново."""
    owner = register(app)
    connect(owner)
    monkeypatch.setattr(bills_bot, "PULL_GAP", 0.0)
    api = FakeBot()
    bills_bot.pull(api, CHAT, NOW)
    line = next(line for line in book["values"][1:] if "№RENT/93" in line and "04.10.2026" in line)
    line[BILL] = "TRUE"
    fresh()
    assert bills_bot.pull(api, CHAT, NOW + timedelta(minutes=1)) == (
        "Списки обновлены: абонентская плата - 2 счёта, аренда - 2 счёта."
    )
    assert api.sent("sendMessage")[-1]["text"].startswith("<b>Аренда")
    assert [item["message_id"] for item in api.sent("deleteMessage")] == [102], "пост аренды заменён, абонентская - на месте"


def press(data: str = "bills:now", chat: dict | None = None) -> dict:
    where = chat or {"id": -1003554193304, "username": "bbc_bills", "type": "channel"}
    return {"update_id": 7, "callback_query": {"id": "q1", "data": data, "message": {"message_id": 5, "chat": where}}}


def test_knopka_cherez_telegram(app: FastAPI, book, monkeypatch) -> None:
    """Нажавшему - «Проверяю книгу…» сразу; итог - сообщением в канале, которое само уходит."""
    owner = register(app)
    connect(owner)
    monkeypatch.setattr(bills_bot, "_now", lambda: NOW)
    later: list[tuple[float, object]] = []
    monkeypatch.setattr(bills_bot, "_later", lambda seconds, action: later.append((seconds, action)))
    api = FakeBot()
    bills_bot.handle(api, CHAT, press())
    assert api.sent("answerCallbackQuery") == [{"callback_query_id": "q1", "text": "Проверяю книгу…"}]
    abon, rent_post, note = api.sent("sendMessage")
    assert [abon["text"], rent_post["text"]] == [ABON_POST, RENT_POST]
    assert note["text"] == f"Списки обновлены: {COUNTS}."
    assert note["disable_notification"] is True and "reply_markup" not in note, "итог - без звука и без кнопки"
    [(seconds, drop)] = later
    assert seconds == bills_bot.NOTE_SECONDS == 5.0
    assert api.sent("deleteMessage") == []
    drop()
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 103}], "итог прочитан - и ушёл"

    # Повторное нажатие в те же 30 секунд - тоже не молчит.
    bills_bot.handle(api, CHAT, press())
    assert api.sent("sendMessage")[-1]["text"] == "Списки только что обновлены - они выше."
    assert len(later) == 2

    bills_bot.handle(api, CHAT, press(chat={"id": -100777, "type": "supergroup"}))
    assert api.sent("answerCallbackQuery")[-1]["text"] == "Кнопка работает только в канале счетов"
    bills_bot.handle(api, CHAT, press("другое"))
    assert api.sent("answerCallbackQuery")[-1]["text"] is None
    assert len(api.sent("sendMessage")) == 4


def test_staroe_nazhatie_ne_otrabatyvaetsya(app: FastAPI, book, monkeypatch) -> None:
    """Нажатие, сделанное, пока бот не слышал канал: Telegram не даёт на него ответить - постов нет."""
    owner = register(app)
    connect(owner)
    monkeypatch.setattr(bills_bot, "_now", lambda: NOW)
    monkeypatch.setattr(bills_bot, "_later", lambda seconds, action: action())
    api = FakeBot()
    api.refuse["answerCallbackQuery"] = TelegramError(
        "Bad Request: query is too old and response timeout expired or query ID is invalid", code=400
    )
    bills_bot.handle(api, CHAT, press())
    assert api.calls == []

    # Ответ не ушёл по сети - нажатие свежее, отрабатывается.
    api.refuse["answerCallbackQuery"] = TelegramError("нет связи с Telegram (ConnectionError)")
    bills_bot.handle(api, CHAT, press())
    assert len(api.sent("sendMessage")) == 3
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 103}]


def test_bind_stavit_knopku(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    api = FakeBot()
    assert bills_bot.ensure_panel(api, CHAT)
    [panel] = api.sent("sendMessage")
    assert panel["text"] == (
        "<b>Счета к выставлению</b>\n"
        "Абонентская плата и аренда - в 10:00, 13:00, 16:00, 19:00 и сразу, когда наступает день счёта."
    )
    assert panel["reply_markup"] == {"inline_keyboard": [[{"text": "Прислать сейчас", "callback_data": "bills:now"}]]}
    assert panel["disable_notification"] is True
    assert api.sent("pinChatMessage") == [{"chat_id": CHAT, "message_id": 101, "disable_notification": True}]

    # Перезапуск: кнопка на месте - правится, вторая не ставится.
    assert bills_bot.ensure_panel(api, CHAT)
    assert len(api.sent("sendMessage")) == 1 and api.sent("editMessageText")[-1]["message_id"] == 101

    # Её удалили из канала - бот при запуске ставит заново.
    api.refuse["editMessageText"] = TelegramError("Bad Request: message to edit not found", code=400)
    assert bills_bot.ensure_panel(api, CHAT)
    del api.refuse["editMessageText"]
    assert [item["message_id"] for item in api.sent("pinChatMessage")] == [101, 102]

    # /bind в канале: команда уходит, кнопка - заново, прежняя откреплена и удалена.
    where = {"id": -1003554193304, "username": "bbc_bills", "type": "channel"}
    bills_bot.handle(api, CHAT, {"update_id": 8, "channel_post": {"message_id": 50, "chat": where, "text": "/bind"}})
    assert api.sent("deleteMessage") == [{"chat_id": CHAT, "message_id": 50}, {"chat_id": CHAT, "message_id": 102}]
    assert api.sent("unpinChatMessage") == [{"chat_id": CHAT, "message_id": 102}]
    assert api.sent("pinChatMessage")[-1]["message_id"] == 103

    # Служебное «закреплено» о своей кнопке - убирается; чужие посты бот не трогает.
    bills_bot.handle(
        api, CHAT, {"update_id": 9, "channel_post": {"message_id": 51, "chat": where, "pinned_message": {"message_id": 103}}}
    )
    bills_bot.handle(api, CHAT, {"update_id": 10, "channel_post": {"message_id": 52, "chat": where, "text": "Привет"}})
    assert api.sent("deleteMessage")[-1] == {"chat_id": CHAT, "message_id": 51}
    assert len(api.sent("deleteMessage")) == 3


# ── Карточка счёта: запись в книгу юрлица ────────────────────────────────────


@pytest.fixture
def books(monkeypatch):
    """Книги юрлиц в памяти: KENESARY живёт в BBC Legal, «Октябрь», строка 6."""
    from bill_book_fake import fake_api

    api = fake_api()
    api.summary_id = "abc123"  # сводка, подключённая в `connect`
    monkeypatch.setattr(bill_book, "_api", api)
    return api


def bill_of(owner: TestClient, customer: str, *, issued: bool = False) -> dict:
    data = owner.get(f"{BASE}/bills").json()
    return next(item for item in data["issued" if issued else "items"] if item["customer"] == customer)


def test_kartochka_pishet_v_knigu_yurlica(app: FastAPI, book, books) -> None:
    owner = register(app)
    connect(owner)
    kenesary = bill_of(owner, "ТОО KENESARY AGRO")

    card = owner.get(f"{BASE}/bills/{kenesary['id']}").json()
    assert card["problem"] == "" and card["can_edit"] is True
    assert card["bill"]["customer"] == "ТОО KENESARY AGRO"
    assert card["book"]["place"]["tab"] == "Октябрь" and card["book"]["place"]["row"] == 6
    shown = card["book"]["fields"]
    assert shown == {"bill_no": "", "bill_date": "", "bill_period": "", "note": "звонить Асель"}

    saved = owner.post(
        f"{BASE}/bills/{kenesary['id']}",
        json={"values": {"bill_no": "312", "bill_date": "2026-10-22"}, "expect": shown},
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["book"]["fields"]["bill_no"] == "312" and saved.json()["book"]["state"] == "TRUE"
    assert saved.json()["changed"] == ["bill_no", "bill_date"]

    # Сводка ещё не догнала книгу, а список уже видит запись.
    data = owner.get(f"{BASE}/bills").json()
    assert all(item["customer"] != "ТОО KENESARY AGRO" for item in data["items"])
    issued = bill_of(owner, "ТОО KENESARY AGRO", issued=True)
    assert (issued["bill_no"], issued["bill_date"]) == ("312", "22.10.2026")
    assert owner.get(f"{BASE}/auth/me").json()["bills"]["soon"] == 2

    # И канал: строка уходит из поста абонентской платы.
    api = FakeBot()
    bills_bot.tick(api, CHAT, NOW)
    assert "KENESARY" not in api.posts("Абонентская плата")[0]["text"]

    # Запись - в журнале KORT: кто, где, что было и что стало.
    from app.finance.models import ActionLog

    with finance_session() as session:
        entry = session.scalars(sa.select(ActionLog).where(ActionLog.kind == "bill.book")).one()
        title, before, after, actor = entry.title, entry.before, entry.after, entry.actor
    assert title == "счёт выставлен в книге: ТОО KENESARY AGRO №ЮО/44, ОКТЯБРЬ 2026 - № 312 от 22.10.2026"
    assert before == {"№ Счета": "", "Дата выстав. Счета": ""} and actor == "owner@bbc.kz"
    assert after["№ Счета"] == "312" and after["вкладка"] == "Октябрь" and after["строка"] == 6

    # Карточка по устаревшим полям - отказ, а не затирание.
    stale = owner.post(f"{BASE}/bills/{kenesary['id']}", json={"values": {"bill_no": "313"}, "expect": shown})
    assert stale.status_code == 409 and "стало «312»" in stale.json()["detail"]


def test_kartochka_bez_knigi_otkryvaetsya(app: FastAPI, book, books) -> None:
    """Книга юрлица не ответила - карточка открывается с причиной, записать нельзя."""
    owner = register(app)
    connect(owner)
    adk = bill_of(owner, "ADK Home")  # его нет в книгах юрлиц
    card = owner.get(f"{BASE}/bills/{adk['id']}").json()
    assert card["book"] is None and "не нашёл ни в одной книге" in card["problem"]
    assert owner.get(f"{BASE}/bills/000000000000").status_code == 404
    assert owner.get(f"{BASE}/bills/not-an-id").status_code == 404


def test_bez_prava_pravki_zapisat_nelzya(app: FastAPI, book, books) -> None:
    owner = register(app)
    connect(owner)
    card = owner.post(
        f"{BASE}/people/employees", json={"full_name": "Юрист Смотрит", "phone": "+77010000002", "access": True}
    )
    assert card.status_code == 201, card.text
    access = owner.put(f"{BASE}/access/employee/{card.json()['id']}", json={"changes": {"invoices": "view"}})
    assert access.status_code == 200, access.text
    lawyer = client(app)
    lawyer.post(f"{BASE}/auth/phone/start", json={"phone": "+77010000002"})
    assert lawyer.post(
        f"{BASE}/auth/phone/set-password", json={"phone": "+77010000002", "password": "secret-123"}
    ).status_code == 200
    kenesary = bill_of(lawyer, "ТОО KENESARY AGRO")
    seen = lawyer.get(f"{BASE}/bills/{kenesary['id']}").json()
    assert seen["can_edit"] is False and seen["book"]["fields"]["bill_no"] == ""
    denied = lawyer.post(f"{BASE}/bills/{kenesary['id']}", json={"values": {"bill_no": "312", "bill_date": "2026-10-22"}})
    assert denied.status_code == 403
    assert books.updates == []


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_shablon_scheta_razdela(app: FastAPI, book, books, monkeypatch) -> None:
    """Шаблон в разделе «Счета»: счета постов одним файлом; заполненный - строками; каждую пишет запись карточки."""
    import io

    from openpyxl import load_workbook

    from app.finance.models import ActionLog

    monkeypatch.setattr(bills_bot, "_now", lambda: NOW)
    owner = register(app)
    connect(owner)
    assert owner.get(f"{BASE}/bills").json()["can_edit"] is True

    got = owner.get(f"{BASE}/bills/template")
    assert got.status_code == 200, got.text
    assert got.headers["content-type"].startswith(XLSX)
    assert "filename*=UTF-8''" in got.headers["content-disposition"] and "22.10.2026" in unquote(got.headers["content-disposition"])
    workbook = load_workbook(io.BytesIO(got.content))
    sheet = workbook.active
    customers = [sheet[f"D{n}"].value for n in range(2, sheet.max_row + 1)]
    # Как в постах: абонентская плата, аренда - по разделам; затем «без вида услуги» и «без даты».
    assert customers == [
        "ТОО KENESARY AGRO", "ИП АЙЯ", "ТОО Find Dream", "ТОО Custom Trade", "ТОО Молфуд", "ADK Home", "ИП ЖАЛМҰХАН",
    ]
    assert workbook.sheetnames == ["Общий реестр", "Аренда", "Абонентские услуги", "BBC", "BBC Omar", "BBCL", "BBCS"]
    assert [workbook["Аренда"][f"D{n}"].value for n in range(2, workbook["Аренда"].max_row + 1)] == [
        "ТОО Find Dream", "ТОО Custom Trade", "ТОО Молфуд", "ИП ЖАЛМҰХАН",
    ]
    assert all(str(sheet[f"N{n}"].value).startswith("=IFERROR(") for n in range(2, sheet.max_row + 1)), "общий реестр тянет"

    # Финансист вписывает на листе своего юрлица.
    def put(title: str, customer: str, number: object, issued: object) -> None:
        firm = workbook[title]
        line = next(n for n in range(2, firm.max_row + 1) if firm[f"D{n}"].value == customer)
        firm[f"N{line}"], firm[f"O{line}"] = number, issued

    put("BBCL", "ТОО KENESARY AGRO", 312, datetime(2026, 10, 22))
    put("BBC Omar", "ТОО Custom Trade", "А-17", "22.10.2026")
    put("BBCS", "ИП АЙЯ", 313, None)  # без даты - отказ этой строке
    out = io.BytesIO()
    workbook.save(out)
    filled = out.getvalue()

    parsed = owner.post(f"{BASE}/bills/template", files={"file": ("шаблон.xlsx", filled, XLSX)})
    assert parsed.status_code == 200, parsed.text
    by_name = {row["customer"]: row for row in parsed.json()["rows"]}
    assert parsed.json()["counts"] == {"ready": 2, "same": 0, "empty": 4, "error": 1}
    assert by_name["ИП АЙЯ"]["problem"] == "«Дата выстав. Счета» - пусто"
    kenesary, trade = by_name["ТОО KENESARY AGRO"], by_name["ТОО Custom Trade"]
    assert kenesary["values"] == {"bill_no": "312", "bill_date": "2026-10-22"}
    assert kenesary["defaults"] == {"bill_period": "с 22.10.2026 года по 21.11.2026 года"}
    # Разбор книги не трогает: счета по-прежнему не выставлены.
    assert books.updates == []
    assert {"ТОО KENESARY AGRO", "ТОО Custom Trade"} <= {item["customer"] for item in owner.get(f"{BASE}/bills").json()["items"]}

    # Каждую готовую строку пишет та же запись, что у карточки - в книгу своего юрлица.
    for row in (kenesary, trade):
        saved = owner.post(
            f"{BASE}/bills/{row['id']}", json={"values": row["values"], "expect": {"bill_no": ""}, "defaults": row["defaults"]}
        )
        assert saved.status_code == 200, saved.text
        assert saved.json()["book"]["state"] == "TRUE"
    legal = books.books["legal-id"].tabs["Октябрь"]
    rent = books.books["rent-id"].tabs["Октябрь"]
    from bill_book_fake import Y, Z

    assert legal.cell(5, Z)["value"] == 312 and legal.cell(5, Y)["value"] == "с 22.10.2026 года по 21.11.2026 года"
    assert rent.cell(3, Z)["value"] == "А-17" and rent.cell(3, Y)["value"] == "с 21.10.2026 года по 20.11.2026 года"
    left = {item["customer"] for item in owner.get(f"{BASE}/bills").json()["items"]}
    assert "ТОО KENESARY AGRO" not in left and "ТОО Custom Trade" not in left, "список видит запись сразу"

    # Тот же файл ещё раз: записанное - «уже в книге», второй записи нет.
    again = owner.post(f"{BASE}/bills/template", files={"file": ("шаблон.xlsx", filled, XLSX)}).json()
    assert {row["customer"]: row["status"] for row in again["rows"] if row["status"] != "empty"} == {
        "ТОО KENESARY AGRO": "same", "ТОО Custom Trade": "same", "ИП АЙЯ": "error",
    }
    # Повтор записи, ответ на которую потерялся: книга уже такая - не конфликт и не вторая запись.
    repeat = owner.post(
        f"{BASE}/bills/{kenesary['id']}", json={"values": kenesary["values"], "expect": {"bill_no": ""}, "defaults": kenesary["defaults"]}
    )
    assert repeat.status_code == 200 and repeat.json()["changed"] == [] and len(books.updates) == 2

    bad = owner.post(f"{BASE}/bills/template", files={"file": ("x.xlsx", b"not a workbook", "text/plain")})
    assert bad.status_code == 422 and "xlsx" in bad.json()["detail"]
    with finance_session() as session:
        logged = [(row.kind, row.category, row.title) for row in session.scalars(sa.select(ActionLog).where(ActionLog.kind == "bill.template"))]
        written = session.scalar(sa.select(sa.func.count()).select_from(ActionLog).where(ActionLog.kind == "bill.book"))
    assert logged == [("bill.template", "export", "шаблон счетов скачан: 7 шт. на 22.10.2026")], "скачивание - в журнале, как выгрузки"
    assert written == 2, "каждая записанная строка - в журнале"

    # Скачать - чтение; загрузить - только с правом «Счета: правка».
    card = owner.post(
        f"{BASE}/people/employees", json={"full_name": "Смотрит Счета", "phone": "+77010000003", "access": True}
    )
    owner.put(f"{BASE}/access/employee/{card.json()['id']}", json={"changes": {"invoices": "view"}})
    viewer = client(app)
    viewer.post(f"{BASE}/auth/phone/start", json={"phone": "+77010000003"})
    viewer.post(f"{BASE}/auth/phone/set-password", json={"phone": "+77010000003", "password": "secret-123"})
    assert viewer.get(f"{BASE}/bills").json()["can_edit"] is False
    assert viewer.get(f"{BASE}/bills/template").status_code == 200
    assert viewer.post(f"{BASE}/bills/template", files={"file": ("x.xlsx", filled, XLSX)}).status_code == 403

    # Выставлять нечего - шаблона нет, сказано словами.
    issue_all(book["values"])
    fresh()
    empty = owner.get(f"{BASE}/bills/template")
    assert empty.status_code == 409 and "нечего" in empty.json()["detail"]


def test_pustoy_list_ne_znachit_vse_vystavleno() -> None:
    with pytest.raises(bills.BillsError, match="ещё собирается"):
        bills.build_board([HEADER], title="", worksheet="Сводка")


def test_nevernyy_chas_vyklyuchaet_tolko_bota(monkeypatch) -> None:
    """«24» в FINANCE_BILLS_HOUR роняло весь API на старте (Render, 22.10.2026)."""
    from app.finance.config import FinanceSettings

    monkeypatch.setenv("FINANCE_BILLS_BOT_TOKEN", "test-token")
    monkeypatch.setenv("FINANCE_BILLS_CHAT_ID", CHAT)
    for raw in ("24", "-1", "десять", ""):
        monkeypatch.setenv("FINANCE_BILLS_HOUR", raw)
        settings = FinanceSettings()
        assert settings.bills_hour_value is None and not settings.bills_ready, raw
    monkeypatch.setenv("FINANCE_BILLS_HOUR", '"9"')
    for raw in ("0", "25", "три"):
        monkeypatch.setenv("FINANCE_BILLS_EVERY", raw)
        assert not FinanceSettings().bills_ready, raw
    monkeypatch.setenv("FINANCE_BILLS_EVERY", "2")
    assert FinanceSettings().bills_hour_value == 9 and FinanceSettings().bills_ready


def test_bez_svodki_bot_molchit(app: FastAPI, book) -> None:
    register(app)
    assert bills_bot.tick(FakeBot(), CHAT, NOW) == "no-workspace"


def test_razdel_ne_vystavleny(app: FastAPI, book) -> None:
    owner = register(app)
    empty = owner.get(f"{BASE}/bills").json()
    assert empty == {"source": None, "items": [], "issued": [], "counts": None, "channel": None, "can_edit": True}
    connect(owner)
    data = owner.get(f"{BASE}/bills").json()
    assert data["counts"]["overdue"] == 8 and data["counts"]["overdue_total"] == "891000"
    assert data["source"]["title"] == "Осн.Общая сводка BBC 2026" and data["source"]["error"] == ""
    assert data["items"][0]["customer"] == "ТОО Молфуд" and data["items"][0]["period_from"] == "2026-06-04"
    assert data["items"][0]["period_to"] == "2026-07-03" and len(data["items"][0]["id"]) == 12
    assert [item["customer"] for item in data["issued"]] == ["Алатау Пауэр"]
    assert data["channel"]["url"] == "https://t.me/bbc_bills" and data["channel"]["hours"] == [10, 13, 16, 19]
    me = owner.get(f"{BASE}/auth/me").json()
    assert me["bills"] == {"overdue": 8, "soon": 3}


def test_razdel_tolko_s_pravom_scheta(app: FastAPI, book) -> None:
    owner = register(app)
    connect(owner)
    card = owner.post(
        f"{BASE}/people/employees", json={"full_name": "Юрист Без Счетов", "phone": "+77010000001", "access": True}
    )
    assert card.status_code == 201, card.text
    lawyer = client(app)
    assert lawyer.post(f"{BASE}/auth/phone/start", json={"phone": "+77010000001"}).json() == {"step": "set_password"}
    done = lawyer.post(f"{BASE}/auth/phone/set-password", json={"phone": "+77010000001", "password": "secret-123"})
    assert done.status_code == 200, done.text
    assert lawyer.get(f"{BASE}/auth/me").json()["access"]["invoices"] == "none"
    assert lawyer.get(f"{BASE}/bills").status_code == 403
    assert lawyer.get(f"{BASE}/auth/me").json()["bills"] is None

    access = owner.put(f"{BASE}/access/employee/{card.json()['id']}", json={"changes": {"invoices": "view"}})
    assert access.status_code == 200, access.text
    assert lawyer.get(f"{BASE}/bills").status_code == 200
    assert lawyer.get(f"{BASE}/auth/me").json()["bills"] == {"overdue": 8, "soon": 3}


def test_kanal_tolko_dlya_svoey_kompanii(app: FastAPI, book, monkeypatch) -> None:
    owner = register(app)
    connect(owner)
    monkeypatch.setattr(finance_settings, "bills_workspace_id", str(uuid.uuid4()))
    assert owner.get(f"{BASE}/bills").json()["channel"] is None
